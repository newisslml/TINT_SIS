"""Lectura de IDs del flujo anterior (cruce por ID_TINT).

El ciclo ya no filtra por ID_TINT (ver adapters/productos.py + sheet_filter.py):
estas funciones quedan para el editor de homologos (cobertura) y para armar la
tabla de productos la primera vez (`cli productos-init`).
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

# El ID de cada formula ("ID_TINT") es un codigo alfanumerico libre
# (LatHab001, EsmSCereluxe Aquatech001, TexSipalinahidr4233, ...). Las hojas de
# homologos traen una columna dedicada con ese encabezado y una fila por ID
# debajo de cada homologo, asi que se lee por encabezado en vez de por patron.
# Se acepta el alias historico "TINT_ID". Ojo: en las hojas de homologos hay
# ademas una columna "ID" (numeros de categoria) que NO es esta, por eso el bare
# "id" no entra aca.
ID_HEADER_ALIASES = {"id_tint", "tint_id"}

# En el experto xData (xData_DATACOMPLETA_*.xlsx) la tabla de formulas esta en la
# hoja "Formulas" (si no existe se usa la primera) y la columna del ID se llama
# "ID" o "ID_TINT" y no siempre es la primera, asi que se ubica por encabezado.
EXPERT_SHEET_NAME = "Formulas"
EXPERT_ID_HEADERS = {"id", "id_tint", "tint_id"}


def _pick_expert_sheet(wb) -> tuple[object, str]:
    for name in wb.sheetnames:
        if name.strip().lower() == EXPERT_SHEET_NAME.lower():
            return wb[name], name
    first = wb.sheetnames[0]
    return wb[first], first


def _find_expert_id_column(header: list) -> int:
    for idx, value in enumerate(header):
        if isinstance(value, str) and value.strip().lower() in EXPERT_ID_HEADERS:
            return idx
    return 0


def read_homologos_ids(path: Path, grupo: str) -> set[str]:
    """Devuelve el set de ID_TINT presentes en la hoja `grupo` del archivo de
    homologos. Ubica la columna por su encabezado ("ID_TINT" / "TINT_ID") y junta
    todos los valores de texto no vacios que aparecen debajo; los renglones de
    estructura (titulo, categorias, encabezados de seccion, filas de homologo)
    dejan esa celda vacia, asi que no se cuelan. Si la hoja no trae esa columna,
    devuelve un set vacio."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[grupo]
        rows = ws.iter_rows(values_only=True)

        id_col: int | None = None
        for row in rows:
            for idx, value in enumerate(row):
                if isinstance(value, str) and value.strip().lower() in ID_HEADER_ALIASES:
                    id_col = idx
                    break
            if id_col is not None:
                break
        if id_col is None:
            return set()

        ids: set[str] = set()
        for row in rows:  # el generador ya quedo posicionado despues del encabezado
            if id_col >= len(row):
                continue
            value = row[id_col]
            if not isinstance(value, str):
                continue
            value = value.strip()
            if value and value.lower() not in ID_HEADER_ALIASES:
                ids.add(value)
        return ids
    finally:
        wb.close()


def read_expert_ids(expert_path: Path) -> set[str]:
    """Devuelve todos los ID_TINT presentes en la hoja "Formulas" del experto
    xData. Se usa para el panel de cobertura del editor de homologos: que IDs del
    experto no estan asignados a ninguna tienda."""
    wb = openpyxl.load_workbook(expert_path, read_only=True, data_only=True)
    try:
        ws, _ = _pick_expert_sheet(wb)
        rows_iter = ws.iter_rows(values_only=True)
        header = list(next(rows_iter))
        id_idx = _find_expert_id_column(header)
        ids: set[str] = set()
        for row in rows_iter:
            if id_idx >= len(row):
                continue
            value = row[id_idx]
            if value is not None:
                text = str(value).strip()
                if text:
                    ids.add(text)
        return ids
    finally:
        wb.close()
