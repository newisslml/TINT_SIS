from __future__ import annotations

import csv
import datetime
from collections.abc import Callable
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import openpyxl

ENCODING = "iso-8859-1"
LINE_TERMINATOR = "\r\n"
DELIMITER = ","
# Confirmado contra TEST/expert.csv (referencia real generada por Corob 4.1.2): las
# columnas de cantidad de colorante ("1/48 onzas") siempre llevan 2 decimales fijos
# (p.ej. "117.00"), mientras que R/G/B/Oz base van como entero plano sin decimales.
ONZAS_HEADER_HINT = "ONZAS"
TWO_DECIMALS = Decimal("0.01")


def _is_onzas_column(header_value: object) -> bool:
    return header_value is not None and ONZAS_HEADER_HINT in str(header_value).upper()


def _format_date(value: datetime.datetime | datetime.date, number_format: str) -> str:
    # Confirmado contra TEST/expert.csv: las celdas de fecha traen 2 number_format
    # distintos en el Excel origen y cada uno se exporta distinto -
    #   'd/m/yyyy;@'  -> sin ceros de relleno, ni en dia ni en mes ("3/8/2026")
    #   cualquier otro (p.ej. 'mm-dd-yy') -> dia y mes con cero de relleno ("04/08/2026")
    # y siempre sin hora (siempre viene en 00:00:00 en el Excel origen).
    if number_format.lower().startswith("d/m/yyyy"):
        return f"{value.day}/{value.month}/{value.year}"
    return f"{value.day:02d}/{value.month:02d}/{value.year}"


def _format_cell(value: object, is_onzas_column: bool, number_format: str) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, (datetime.datetime, datetime.date)):
        return _format_date(value, number_format)
    if isinstance(value, (int, float)):
        if is_onzas_column:
            # Redondeo "half up" (3.125 -> 3.13), no el "half to even" que usa el
            # formato %.2f de Python (3.125 -> 3.12 por defecto de IEEE 754).
            return str(Decimal(str(value)).quantize(TWO_DECIMALS, rounding=ROUND_HALF_UP))
        if float(value) == int(value):
            return str(int(value))
        return str(value)
    return str(value)


def write_passthrough_csv(
    input_path: Path,
    output_path: Path,
    *,
    on_progress: Callable[[dict], None] | None = None,
    total_hint: int | None = None,
    progress_every: int = 2000,
) -> int:
    """Convierte un Excel que ya viene en formato final (una fila por formula, con
    Clasificacion/Producto/Cartilla/Formato/... como columnas propias) directo a
    CSV, sin filtrar, limpiar ni validar nada: tintometria ya entrega la tabla
    completa lista para la maquina, aca solo cambia el contenedor (.xlsx -> .csv).

    Devuelve la cantidad de filas de datos escritas (sin contar el encabezado).

    `on_progress`, si se pasa, recibe {leidas, total} cada `progress_every` filas;
    `total_hint` es el total esperado (lo sabe quien llama tras filtrar).
    """
    input_path = Path(input_path)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(input_path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    total = total_hint
    if total is None:
        mr = ws.max_row
        total = mr - 1 if isinstance(mr, int) and mr > 0 else None
    rows_iter = ws.iter_rows()

    header = [c.value for c in next(rows_iter)]
    onzas_flags = [_is_onzas_column(h) for h in header]

    rows_written = 0
    with open(output_path, "w", newline="", encoding=ENCODING) as f:
        writer = csv.writer(f, delimiter=DELIMITER, lineterminator=LINE_TERMINATOR)
        writer.writerow(["" if h is None else str(h) for h in header])
        for row in rows_iter:
            writer.writerow(
                _format_cell(
                    cell.value,
                    onzas_flags[i] if i < len(onzas_flags) else False,
                    cell.number_format,
                )
                for i, cell in enumerate(row)
            )
            rows_written += 1
            if on_progress is not None and rows_written % progress_every == 0:
                on_progress({"leidas": rows_written, "total": total})

    if on_progress is not None:
        on_progress({"leidas": rows_written, "total": total or rows_written})

    wb.close()
    return rows_written
