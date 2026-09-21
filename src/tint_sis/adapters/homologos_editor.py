"""Parser y editor de `homologos_TINT.xlsx` para la vista /homologos de la app.

El archivo real (confirmado corriendo esto contra `data/input/homologos_TINT.xlsx`,
4 hojas de hasta ~186k filas) sigue, por hoja, esta estructura de bloques:

  fila 0        titulo ("LISTADO PRODUCTOS HOMOLOGOS ...")
  filas 1-2     en blanco
  fila 3        encabezado de la tabla de lineas: PRODUCTOS | ID_TINT | ID | CODE | DESCR | PATH
  filas 4..N    una fila por linea (catalogo): C=ID numerico, D=CODE, E=DESCR (==D), F=PROD####
  (blanco)
  por cada linea, una "seccion":
    fila seccion   A="<nombre> -  <PROD####>", C/D/E/F = strings meta tipo "ID,C,5" (ancho de columna
                   del formato original; no lo usa el resto del sistema, se preserva tal cual)
    por cada homologo de esa linea:
      fila homologo   A=None (o una nota suelta, ver mas abajo), B=None, C=ID numerico o None,
                      D=CODE, E=DESCR (==D), F=SUBP#### o None
      filas ID_TINT   A=None, B=<ID_TINT>, D=CODE (repetido) -- una por cada ID asignado

Un homologo sin ninguna fila de ID_TINT debajo esta "pendiente" (existe la linea/producto
pero todavia no tiene formulas asignadas) -- ver VISTAS_TINT_SIS.md #8.

El archivo real trae ademas un puñado de filas que no calzan del todo con el patron
(un homologo con el ID numerico vacio, un homologo marcado con "????" en la columna A
en vez de dejarla vacia, una fila "ID" suelta). El parser es tolerante: lo que no
reconoce lo cuenta en `filas_no_reconocidas` en vez de romper, y no lo toca al guardar
(solo se insertan/borran filas en los lugares que el editor conoce).
"""
from __future__ import annotations

import re
import shutil
from copy import copy as _copy
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

LINEA_HEADER = "PRODUCTOS"
_SECCION_RE = re.compile(r"^(?P<nombre>.*?)\s*-\s*(?P<path>[A-Za-z]+\d+)\s*$")


@dataclass
class Homologo:
    fila: int
    id_num: int | None
    nombre: str
    path: str | None
    nota: str | None = None
    id_rows: list[tuple[int, str]] = field(default_factory=list)

    @property
    def ids(self) -> list[str]:
        return [v for _, v in self.id_rows]

    @property
    def pendiente(self) -> bool:
        return not self.id_rows


@dataclass
class Linea:
    nombre: str
    id_num: int | None
    code: str | None
    descr: str | None
    path: str | None
    fila_tabla: int | None = None
    fila_seccion: int | None = None
    fila_seccion_fin: int = 0
    meta: tuple | None = None
    homologos: list[Homologo] = field(default_factory=list)

    @property
    def pendiente(self) -> bool:
        return not self.homologos


@dataclass
class HomologosSheet:
    titulo: str | None
    lineas: list[Linea]
    filas_no_reconocidas: int
    ultima_fila: int


def _split_seccion(valor: object) -> tuple[str, str | None]:
    if not isinstance(valor, str):
        return (str(valor), None)
    m = _SECCION_RE.match(valor.strip())
    if m:
        return (m.group("nombre").strip(), m.group("path"))
    return (valor.strip(), None)


def _linea_por_path(lineas: list[Linea], path: str | None) -> Linea | None:
    if not path:
        return None
    for linea in lineas:
        if linea.path == path:
            return linea
    return None


def _clasificar_hoja(ws) -> HomologosSheet:
    titulo: str | None = None
    lineas: list[Linea] = []
    current_linea: Linea | None = None
    current_homologo: Homologo | None = None
    in_sections = False
    no_reconocidas = 0
    ultima_fila = 0

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        fila = i + 1
        vals = (tuple(row) + (None,) * 6)[:6]
        a, b, c, d, e, f = vals
        if all(v is None for v in vals):
            continue
        ultima_fila = fila

        if i == 0 and b is None and c is None and isinstance(a, str):
            titulo = a
            continue
        if a == LINEA_HEADER:
            continue
        if isinstance(c, str) and "," in c:
            nombre_seccion, path_seccion = _split_seccion(a)
            linea = _linea_por_path(lineas, path_seccion)
            if linea is None:
                linea = Linea(
                    nombre=nombre_seccion,
                    id_num=None,
                    code=None,
                    descr=None,
                    path=path_seccion,
                )
                lineas.append(linea)
            else:
                linea.nombre = nombre_seccion
            linea.fila_seccion = fila
            linea.fila_seccion_fin = fila
            linea.meta = (c, d, e, f)
            current_linea = linea
            current_homologo = None
            in_sections = True
            continue
        # CODE (d) y DESCR (e) deberian ser iguales pero el archivo real trae
        # typos/abreviaciones donde difieren (confirmado: "Esm. Sintéti" vs
        # "Esm. Sintéticos", "Invictus Semi Satin" vs "... Ceresita"); exigir
        # d == e dejaba esas filas sin reconocer y sus ID_TINT de abajo se
        # colaban en el homologo anterior. La posicion (antes/despues de la
        # primera seccion) y que b este vacio ya alcanzan para distinguirlas
        # de una fila de ID_TINT (que siempre trae b) o de un encabezado de
        # seccion (ya filtrado arriba).
        if not in_sections and b is None and d is not None:
            path_val = f if isinstance(f, str) else None
            linea = _linea_por_path(lineas, path_val)
            if linea is None:
                linea = Linea(
                    nombre=str(d),
                    id_num=c if isinstance(c, int) else None,
                    code=str(d),
                    descr=str(e) if e is not None else str(d),
                    path=path_val,
                    fila_tabla=fila,
                )
                lineas.append(linea)
            else:
                linea.fila_tabla = fila
                linea.id_num = c if isinstance(c, int) else linea.id_num
            continue
        if in_sections and b is None and d is not None:
            if current_linea is None:
                no_reconocidas += 1
                continue
            homologo = Homologo(
                fila=fila,
                id_num=c if isinstance(c, int) else None,
                nombre=str(d),
                path=f if isinstance(f, str) else None,
                nota=a if isinstance(a, str) else None,
            )
            current_linea.homologos.append(homologo)
            current_linea.fila_seccion_fin = fila
            current_homologo = homologo
            continue
        if in_sections and isinstance(b, str) and b.strip():
            if current_homologo is None:
                no_reconocidas += 1
                continue
            current_homologo.id_rows.append((fila, b.strip()))
            if current_linea is not None:
                current_linea.fila_seccion_fin = fila
            continue
        no_reconocidas += 1

    return HomologosSheet(
        titulo=titulo, lineas=lineas, filas_no_reconocidas=no_reconocidas, ultima_fila=ultima_fila
    )


def list_sheet_names(path: Path) -> list[str]:
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def parse_sheet(path: Path, sheet_name: str) -> HomologosSheet:
    """Lectura rapida (read-only) de una hoja, sin poder editarla despues.
    Se usa para paneles que solo necesitan mostrar datos (ej. cobertura desde
    otra pantalla) sin pagar el costo de cargar el workbook completo en modo
    editable."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        return _clasificar_hoja(wb[sheet_name])
    finally:
        wb.close()


def compute_cobertura(sheet: HomologosSheet, expert_ids: set[str]) -> dict:
    asignados: dict[str, list[tuple[str, str]]] = {}
    homologos_sin_ids = []
    for linea in sheet.lineas:
        for homologo in linea.homologos:
            if not homologo.ids:
                homologos_sin_ids.append({"linea": linea.nombre, "homologo": homologo.nombre})
            for id_tint in homologo.ids:
                asignados.setdefault(id_tint, []).append((linea.nombre, homologo.nombre))

    todos_asignados = set(asignados)
    duplicados = [
        {
            "id_tint": id_tint,
            "asignaciones": [{"linea": ln, "homologo": hn} for ln, hn in usos],
        }
        for id_tint, usos in asignados.items()
        if len(usos) > 1
    ]
    return {
        "ids_sin_asignar": sorted(expert_ids - todos_asignados),
        "ids_no_en_experto": sorted(todos_asignados - expert_ids),
        "ids_duplicados": duplicados,
        "homologos_sin_ids": homologos_sin_ids,
    }


def _copiar_estilo_fila(ws, origen: int, destino: int, ncols: int = 6) -> None:
    for col in range(1, ncols + 1):
        src = ws.cell(row=origen, column=col)
        dst = ws.cell(row=destino, column=col)
        dst.font = _copy(src.font)
        dst.fill = _copy(src.fill)
        dst.border = _copy(src.border)
        dst.alignment = _copy(src.alignment)
        dst.number_format = src.number_format


def _fila_encabezado_tabla(ws) -> int:
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True)):
        if row and row[0] == LINEA_HEADER:
            return i + 1
    raise ValueError("No se encontro la fila de encabezado de la tabla de lineas ('PRODUCTOS')")


def _buscar_linea(sheet: HomologosSheet, nombre: str) -> Linea:
    for linea in sheet.lineas:
        if linea.nombre == nombre:
            return linea
    raise ValueError(f"Linea '{nombre}' no encontrada")


def _buscar_homologo(linea: Linea, nombre: str) -> Homologo:
    for homologo in linea.homologos:
        if homologo.nombre == nombre:
            return homologo
    raise ValueError(f"Homologo '{nombre}' no encontrado en la linea '{linea.nombre}'")


class HomologosEditor:
    """Envuelve el workbook completo (las 4 hojas) abierto en modo editable.

    Cargarlo tarda (~decenas de segundos: el archivo real tiene ~700k celdas
    entre las 4 hojas), asi que se abre una sola vez y se mantiene en memoria
    entre pedidos de la UI (ver `_get_editor` en api.py); los cambios quedan
    en memoria hasta llamar a `guardar()`, que ademas deja una copia de
    respaldo (`<nombre>_backup.xlsx`) de la version anterior por si algo sale
    mal.
    """

    def __init__(self, path: Path):
        self._path = Path(path)
        self._wb = openpyxl.load_workbook(self._path)
        self._cache: dict[str, HomologosSheet] = {}
        self._dirty = False

    @property
    def path(self) -> Path:
        return self._path

    @property
    def dirty(self) -> bool:
        return self._dirty

    def tiendas(self) -> list[str]:
        return list(self._wb.sheetnames)

    def arbol(self, tienda: str) -> HomologosSheet:
        if tienda not in self._cache:
            self._cache[tienda] = _clasificar_hoja(self._wb[tienda])
        return self._cache[tienda]

    def _invalidar(self, tienda: str) -> None:
        self._cache.pop(tienda, None)
        self._dirty = True

    def agregar_id(self, tienda: str, linea_nombre: str, homologo_nombre: str, id_tint: str) -> None:
        id_tint = (id_tint or "").strip()
        if not id_tint:
            raise ValueError("El ID_TINT no puede estar vacio")
        ws = self._wb[tienda]
        sheet = self.arbol(tienda)
        linea = _buscar_linea(sheet, linea_nombre)
        homologo = _buscar_homologo(linea, homologo_nombre)
        if id_tint in homologo.ids:
            raise ValueError(f"'{id_tint}' ya esta asignado a '{homologo.nombre}'")

        if homologo.id_rows:
            estilo_desde = homologo.id_rows[-1][0]
        else:
            estilo_desde = homologo.fila
        insertar_en = estilo_desde + 1

        ws.insert_rows(insertar_en)
        _copiar_estilo_fila(ws, estilo_desde, insertar_en)
        ws.cell(row=insertar_en, column=2, value=id_tint)
        ws.cell(row=insertar_en, column=4, value=homologo.nombre)
        self._invalidar(tienda)

    def quitar_id(self, tienda: str, linea_nombre: str, homologo_nombre: str, id_tint: str) -> None:
        ws = self._wb[tienda]
        sheet = self.arbol(tienda)
        linea = _buscar_linea(sheet, linea_nombre)
        homologo = _buscar_homologo(linea, homologo_nombre)
        fila = next((f for f, v in homologo.id_rows if v == id_tint), None)
        if fila is None:
            raise ValueError(f"'{id_tint}' no esta asignado a '{homologo.nombre}'")
        ws.delete_rows(fila)
        self._invalidar(tienda)

    def agregar_homologo(self, tienda: str, linea_nombre: str, nombre: str) -> None:
        nombre = (nombre or "").strip()
        if not nombre:
            raise ValueError("El nombre del homologo no puede estar vacio")
        ws = self._wb[tienda]
        sheet = self.arbol(tienda)
        linea = _buscar_linea(sheet, linea_nombre)
        if any(h.nombre == nombre for h in linea.homologos):
            raise ValueError(f"Ya existe un homologo '{nombre}' en '{linea.nombre}'")
        if linea.fila_seccion is None:
            raise ValueError(f"La linea '{linea.nombre}' todavia no tiene seccion propia en la hoja")

        ids_existentes = [h.id_num for h in linea.homologos if h.id_num is not None]
        nuevo_id = (max(ids_existentes) + 1) if ids_existentes else 1
        nuevo_path = f"SUBP{nuevo_id:04d}"
        insertar_en = linea.fila_seccion_fin + 1

        ws.insert_rows(insertar_en)
        if linea.homologos:
            _copiar_estilo_fila(ws, linea.homologos[-1].fila, insertar_en)
        ws.cell(row=insertar_en, column=3, value=nuevo_id)
        ws.cell(row=insertar_en, column=4, value=nombre)
        ws.cell(row=insertar_en, column=5, value=nombre)
        ws.cell(row=insertar_en, column=6, value=nuevo_path)
        self._invalidar(tienda)

    def agregar_linea(self, tienda: str, nombre: str) -> None:
        nombre = (nombre or "").strip()
        if not nombre:
            raise ValueError("El nombre de la linea no puede estar vacio")
        ws = self._wb[tienda]
        sheet = self.arbol(tienda)
        if any(l.nombre == nombre for l in sheet.lineas):
            raise ValueError(f"Ya existe una linea '{nombre}'")

        ids_tabla = [l.id_num for l in sheet.lineas if l.id_num is not None]
        nuevo_id = (max(ids_tabla) + 1) if ids_tabla else 1
        nuevo_path = f"PROD{nuevo_id:04d}"

        con_tabla = [l for l in sheet.lineas if l.fila_tabla is not None]
        if con_tabla:
            estilo_tabla_desde = max(l.fila_tabla for l in con_tabla)
            insertar_tabla_en = estilo_tabla_desde + 1
        else:
            estilo_tabla_desde = None
            insertar_tabla_en = _fila_encabezado_tabla(ws) + 1

        ws.insert_rows(insertar_tabla_en)
        if estilo_tabla_desde is not None:
            _copiar_estilo_fila(ws, estilo_tabla_desde, insertar_tabla_en)
        ws.cell(row=insertar_tabla_en, column=3, value=nuevo_id)
        ws.cell(row=insertar_tabla_en, column=4, value=nombre)
        ws.cell(row=insertar_tabla_en, column=5, value=nombre)
        ws.cell(row=insertar_tabla_en, column=6, value=nuevo_path)
        self._invalidar(tienda)

        sheet = self.arbol(tienda)
        meta = sheet.lineas[-1].meta if sheet.lineas and sheet.lineas[-1].meta else (
            "ID,C,5",
            "CODE,C,40",
            "DESCR,C,40",
            "PATH,C,12",
        )
        estilo_seccion_desde = sheet.lineas[-1].fila_seccion if sheet.lineas else None
        fila_seccion = sheet.ultima_fila + 2

        ws.cell(row=fila_seccion, column=1, value=f"{nombre} -  {nuevo_path}")
        ws.cell(row=fila_seccion, column=3, value=meta[0])
        ws.cell(row=fila_seccion, column=4, value=meta[1])
        ws.cell(row=fila_seccion, column=5, value=meta[2])
        ws.cell(row=fila_seccion, column=6, value=meta[3])
        if estilo_seccion_desde is not None:
            _copiar_estilo_fila(ws, estilo_seccion_desde, fila_seccion)
        self._invalidar(tienda)

    def guardar(self) -> Path:
        # Extension .xlsx (no ".bak") a proposito: asi el respaldo se puede
        # abrir directo con Excel/openpyxl si algo sale mal.
        backup = self._path.with_name(f"{self._path.stem}_backup{self._path.suffix}")
        shutil.copy2(self._path, backup)
        self._wb.save(self._path)
        self._dirty = False
        return backup
