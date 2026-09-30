"""Preparacion de los archivos expertos de un ciclo a partir del archivo maestro
que envia tintometria (vista "Preparar experto" y `cli preparar`).

1. `analizar_maestro` lee el maestro (formato Experto 1, ver maestro.py) y lo
   compara con:
     - el Experto 1 del ciclo anterior: productos nuevos, productos que ya no
       vienen y formulas agregadas/quitadas por producto;
     - la tabla de productos: productos que no estan (se agregan al preparar);
     - los Experto 2/3 anteriores (plantillas): que productos entran y salen.
   Tambien valida que se pueda convertir (formatos con equivalencia a galon,
   colorantes legibles) y cuenta las formulas duplicadas: filas con la misma
   Clasificacion, Producto, Cartilla, Color y Base (maestro.CLAVE_FORMULA) que
   una anterior, aunque cambien RGB, notas, formato o colorantes. Se deja la
   primera (Texturex llego con sus 4.775 formulas copiadas 6 veces; Construcolor
   con 4.064 colores dos veces). Experto 3 no lleva cartilla: ahi ademas se deja
   una sola vez el mismo color en distintas cartillas (maestro.CLAVE_FORMULA_E3).
2. `preparar_expertos`, con lo que decide el usuario para cada producto nuevo
   (tiendas y nombres):
     - agrega esos productos a productos_TINT.xlsx y completa los nombres de
       Experto 2/3 que falten (todos los productos del maestro van a los 3);
     - genera Experto_1/2/3_<DD_MM_YYYY> en la carpeta de entrada: E1 es el
       maestro (sin las duplicadas); E2 y E3 salen de las reglas de maestro.py
       sobre el libro del ciclo anterior como plantilla (macros, hojas
       auxiliares e IntegrityData se conservan);
     - mueve los expertos y la tabla anteriores a <data>/backups/expertos/<fecha>/;
     - deja un resumen (Excel + JSON) en <salida>/Preparacion expertos/.
"""
from __future__ import annotations

import datetime
import json
import os
import re
import shutil
import zipfile
from collections import Counter
from collections.abc import Callable, Iterator, Mapping
from dataclasses import asdict, dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from tint_sis import maestro
from tint_sis.adapters.libro_writer import encabezado as encabezado_libro
from tint_sis.adapters.libro_writer import reescribir_hoja
from tint_sis.adapters.productos import FilaTabla, ProductoAgregado, actualizar_tabla, leer_filas_tabla
from tint_sis.adapters.sheet_filter import FormatoExpertoError, contar_valores, quitar_filas_repetidas
from tint_sis.backups import FORMATO_CARPETA, carpeta_backups
from tint_sis.config import AppConfig
from tint_sis.expertos import EXPERTOS, KEY_SEP, normalizar
from tint_sis.routing import find_latest_expert, matching_files

ProgressCallback = Callable[[dict], None]

E1, E2, E3 = "Experto 1", "Experto 2", "Experto 3"
PREPARACION_DIRNAME = "Preparacion expertos"
BACKUP_EXPERTOS_DIRNAME = "expertos"
STAGING_DIRNAME = ".preparando"
# "Todo  MP14 28092026", "..._28_09_2026", "28-09-2026": la ultima fecha del nombre
_FECHA_RE = re.compile(r"(\d{2})[ _.-]?(\d{2})[ _.-]?(\d{4})(?!.*\d)")
_MAX_EJEMPLOS = 5


class PreparacionError(Exception):
    """No se pudo preparar (archivo abierto, falta algo). Se avisa sin traceback."""


def _emit(cb: ProgressCallback | None, **evento: object) -> None:
    if cb is not None:
        cb(evento)


def _miles(n: int) -> str:
    return f"{n:,}".replace(",", ".")


def fecha_maestro(path: Path, hoy: datetime.date | None = None) -> tuple[str, bool]:
    """Fecha DD_MM_YYYY del maestro: la de su nombre ("Todo  MP14 28092026") o,
    si no la trae, la de hoy. El bool dice si vino del nombre."""
    m = _FECHA_RE.search(Path(path).stem)
    if m:
        d, mes, a = (int(x) for x in m.groups())
        try:
            return datetime.date(a, mes, d).strftime("%d_%m_%Y"), True
        except ValueError:
            pass
    return (hoy or datetime.date.today()).strftime("%d_%m_%Y"), False


def nombre_experto(glob: str, fecha: str, ext: str) -> str:
    """Experto_1*.xlsx + 28_09_2026 -> Experto_1_28_09_2026.xlsx"""
    prefijo = glob.split("*", 1)[0].rstrip("_")
    return f"{prefijo}_{fecha}{ext}"


def _canon(conteo: Mapping[object, int]) -> dict[str, str]:
    """valor normalizado -> la grafia mas usada."""
    out: dict[str, str] = {}
    for valor, _ in sorted(conteo.items(), key=lambda kv: -kv[1]):
        if valor is not None and str(valor).strip():
            out.setdefault(normalizar(valor), str(valor))
    return out


def _crudo(fila: list) -> tuple:
    """Contenido de una fila para comparar con el experto anterior (sin celdas
    vacias al final)."""
    crudo = tuple(fila)
    while crudo and crudo[-1] in (None, ""):
        crudo = crudo[:-1]
    return crudo


def _nombre_e3(nombre: str, pares: Mapping[str, tuple[str, str]], grupos: Mapping[str, str]) -> tuple[str, str]:
    """"Látex / Habitacional" (tabla) -> (group_code, product_code) con la grafia
    exacta de la plantilla ("Látex " lleva un espacio al final). Sin separador,
    el grupo queda vacio."""
    par = pares.get(normalizar(nombre))
    if par is not None:
        return par
    sep = KEY_SEP if KEY_SEP in nombre else "/"
    grupo, hay, producto = nombre.partition(sep)
    if not hay:
        return "", nombre.strip()
    grupo = grupo.strip()
    return grupos.get(normalizar(grupo), grupo), producto.strip()


# --------------------------------------------------------------------------- #
# modelo
# --------------------------------------------------------------------------- #
@dataclass
class ProductoMaestro:
    producto: str  # nombre visible (sin espacios de los bordes)
    clasificacion: str = ""
    filas: int = 0
    repetidas: int = 0  # formulas duplicadas (misma maestro.CLAVE_FORMULA que una anterior)
    distintas: int = 0  # de esas, con colorantes distintos a los de la primera
    # ademas, solo en Experto 3: mismo color y base que una anterior en otra cartilla
    repetidas_e3: int = 0
    distintas_e3: int = 0
    formatos: Counter = field(default_factory=Counter)
    cartillas: Counter = field(default_factory=Counter)
    clasificaciones: Counter = field(default_factory=Counter)

    @property
    def unicas(self) -> int:
        return self.filas - self.repetidas


@dataclass
class ProductoNuevo:
    """Producto del maestro que no esta en la tabla de productos."""

    clave: str
    producto: str
    clasificacion: str
    linea: str
    filas: int
    repetidas: int
    formatos: dict[str, int]
    cartillas: dict[str, int]
    nombre_e2: str
    nombre_e3: str
    tiendas: list[str]
    nuevo_en_maestro: bool  # tampoco venia en el Experto 1 anterior


@dataclass
class Cambio:
    producto: str
    antes: int
    ahora: int
    agregadas: int
    quitadas: int


@dataclass
class PlanExperto:
    label: str
    archivo: str | None
    plantilla: str | None
    productos: int
    filas: int
    entran: list[str]  # productos que el experto anterior no tenia
    salen: list[str]  # productos del experto anterior que el maestro ya no trae
    estado: str  # "ok" | "sin-plantilla"


@dataclass
class Analisis:
    maestro: str
    ruta_maestro: str
    fecha: str
    fecha_del_nombre: bool
    filas: int
    filas_unicas: int
    repetidas: int
    productos: int
    anterior: str | None
    anterior_filas: int
    anterior_productos: int
    nuevos: list[ProductoNuevo]
    quitados: list[dict]
    cambios: list[Cambio]
    sin_cambios: int
    filas_agregadas: int
    filas_quitadas: int
    tabla_sin_maestro: list[str]
    expertos: list[PlanExperto]
    formatos: dict[str, dict]
    tiendas: list[str]
    advertencias: list[str]
    bloqueantes: list[str]
    puede_preparar: bool = False
    repetidas_e3: int = 0  # ademas de `repetidas`, las que no van a Experto 3 (sin cartilla)
    # lo que necesita preparar_expertos (no va a la UI)
    _nombres: dict = field(default_factory=dict, repr=False)  # clave -> maestro.NombresProducto
    _completar: dict = field(default_factory=dict, repr=False)  # fila de la tabla -> experto -> nombre
    _plantillas: dict = field(default_factory=dict, repr=False)  # label -> Path
    _pares_e3: dict = field(default_factory=dict, repr=False)
    _grupos_e3: dict = field(default_factory=dict, repr=False)
    _bases_e3: dict = field(default_factory=dict, repr=False)
    _clasif_e2: dict = field(default_factory=dict, repr=False)
    _productos_e2: dict = field(default_factory=dict, repr=False)
    _grupo_de: dict = field(default_factory=dict, repr=False)  # clasificacion normalizada -> grupo E3
    _duplicados: list = field(default_factory=list, repr=False)  # Duplicada de cada fila que se quita

    def to_dict(self) -> dict:
        return {k: v for k, v in asdict(self).items() if not k.startswith("_")}


@dataclass
class Duplicada:
    """Fila del maestro que se quita por repetir la formula de una anterior."""

    producto: str
    clasificacion: str
    cartilla: str
    color: str
    base: str
    fila_conservada: int  # fila de Excel del maestro
    fila_quitada: int
    misma_formula: bool  # mismos colorantes (por galon) que la conservada
    solo_e3: bool = False  # otra cartilla: se quita solo de Experto 3
    cartilla_conservada: str = ""


@dataclass
class _Lectura:
    productos: dict[str, ProductoMaestro]
    contenido: dict[str, Counter]  # clave -> contenido de cada fila (sin duplicadas) -> veces
    formatos: dict[str, dict]
    ilegibles: int
    ejemplos_ilegibles: list[str]
    max_colorantes: int
    duplicados: list[Duplicada] = field(default_factory=list)


def _colorantes(cols: maestro.Columnas, fila: list) -> tuple:
    """Colorantes de la fila por galon, para ver si dos duplicadas traen la misma
    formula (texto crudo si alguno no se puede leer o el formato no se conoce)."""
    celdas = cols.celdas_colorante(fila)
    galones = maestro.galones(cols.valor(fila, "formato"))
    if galones is not None:
        try:
            return tuple(sorted(maestro.colorantes_por_galon(celdas, galones)))
        except maestro.ColoranteError:
            pass
    return tuple(normalizar(c) for c in celdas)


# --------------------------------------------------------------------------- #
# analisis
# --------------------------------------------------------------------------- #
def _leer(path: Path, on_progress: ProgressCallback | None, etapa: str, pct: tuple[int, int]) -> _Lectura:
    """Recorre un maestro / Experto 1: productos, contenido de las filas por
    producto (para comparar), formatos, colorantes ilegibles y formulas
    duplicadas (misma maestro.CLAVE_FORMULA que una fila anterior)."""
    lo, hi = pct

    def progreso(ev: dict) -> None:
        total = ev.get("total") or 0
        frac = min(1.0, (ev.get("leidas") or 0) / total) if total else 0.0
        _emit(on_progress, fase="progreso", etapa=etapa, pct=lo + int(frac * (hi - lo)),
              leidas=ev.get("leidas"), total=total)

    cols, filas = maestro.leer(path, on_progress=progreso if on_progress else None)
    lectura = _Lectura({}, {}, {}, 0, [], 0)
    # clave de formula -> (fila, colorantes) de la primera; la de Experto 3 sin
    # cartilla, con la cartilla de la primera
    vistas: dict[tuple, tuple[int, tuple]] = {}
    vistas_e3: dict[tuple, tuple[int, tuple, str]] = {}
    for n, fila in enumerate(filas, start=2):
        nombre = cols.valor(fila, "producto")
        if nombre is None or not str(nombre).strip():
            continue
        clave = normalizar(nombre)
        p = lectura.productos.get(clave)
        if p is None:
            p = lectura.productos[clave] = ProductoMaestro(producto=str(nombre).strip())
            lectura.contenido[clave] = Counter()
        p.filas += 1
        p.clasificaciones[str(cols.valor(fila, "clasificacion") or "").strip()] += 1
        formato = str(cols.valor(fila, "formato") or "").strip()
        p.formatos[formato] += 1
        p.cartillas[str(cols.valor(fila, "cartilla") or "").strip()] += 1
        cartilla = str(cols.valor(fila, "cartilla") or "").strip()
        colorantes = _colorantes(cols, fila)
        formula = cols.clave_formula(fila)
        formula_e3 = cols.clave_formula(fila, maestro.CLAVE_FORMULA_E3)
        primera = vistas.get(formula)
        primera_e3 = vistas_e3.get(formula_e3)
        solo_e3 = False
        if primera is None:
            vistas[formula] = (n, colorantes)
            lectura.contenido[clave][_crudo(fila)] += 1  # solo lo que se entrega
            if primera_e3 is None:
                vistas_e3[formula_e3] = (n, colorantes, cartilla)
            else:  # el mismo color en otra cartilla: no va a Experto 3
                solo_e3 = True
                primera = primera_e3[:2]
        if primera is not None:
            misma = primera[1] == colorantes
            if solo_e3:
                p.repetidas_e3 += 1
                p.distintas_e3 += not misma
            else:
                p.repetidas += 1
                p.distintas += not misma
            lectura.duplicados.append(
                Duplicada(
                    producto=p.producto,
                    clasificacion=str(cols.valor(fila, "clasificacion") or "").strip(),
                    cartilla=cartilla,
                    color=str(cols.valor(fila, "color") or "").strip(),
                    base=str(cols.valor(fila, "base") or "").strip(),
                    fila_conservada=primera[0],
                    fila_quitada=n,
                    misma_formula=misma,
                    solo_e3=solo_e3,
                    cartilla_conservada=primera_e3[2] if solo_e3 else cartilla,
                )
            )
        f = lectura.formatos.setdefault(formato, {"filas": 0, "galones": maestro.galones(formato)})
        f["filas"] += 1
        usados = 0
        for celda in cols.celdas_colorante(fila):
            try:
                if maestro.parse_colorante(celda) is not None:
                    usados += 1
            except maestro.ColoranteError:
                lectura.ilegibles += 1
                if len(lectura.ejemplos_ilegibles) < _MAX_EJEMPLOS:
                    lectura.ejemplos_ilegibles.append(f"fila {n}: '{celda}'")
        lectura.max_colorantes = max(lectura.max_colorantes, usados)
    for p in lectura.productos.values():
        p.clasificacion = p.clasificaciones.most_common(1)[0][0]
    return lectura


def analizar_maestro(path: Path, cfg: AppConfig, *, on_progress: ProgressCallback | None = None) -> Analisis:
    """Compara el maestro con el ciclo anterior y la tabla de productos, y arma
    el plan de los 3 expertos. No escribe nada."""
    path = Path(path)
    input_dir = Path(cfg.input_dir)
    advertencias: list[str] = []
    bloqueantes: list[str] = []

    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        raise FormatoExpertoError(f"{path.name}: el maestro tiene que venir en .xlsx (el .xls corta en 65.535 filas)")
    tabla_path = input_dir / cfg.productos_name
    if not tabla_path.exists():
        raise FormatoExpertoError(f"Falta la tabla de productos ({cfg.productos_name}) en la carpeta de entrada")
    fecha, del_nombre = fecha_maestro(path)
    if not del_nombre:
        advertencias.append(
            f"{path.name} no trae la fecha en el nombre (DDMMAAAA): los expertos salen con la de hoy, {fecha}"
        )

    # --- maestro ---
    _emit(on_progress, fase="etapa", etapa="Leyendo el maestro", pct=1)
    lec = _leer(path, on_progress, "Leyendo el maestro", (1, 40))
    productos = lec.productos
    if not productos:
        raise FormatoExpertoError(f"{path.name} no tiene formulas")
    filas = sum(p.filas for p in productos.values())
    repetidas = sum(p.repetidas for p in productos.values())
    unicas = filas - repetidas
    repetidas_e3 = sum(p.repetidas_e3 for p in productos.values())

    sin_equivalencia = [f"'{f}' ({_miles(d['filas'])} formulas)" for f, d in lec.formatos.items() if d["galones"] is None]
    if sin_equivalencia:
        bloqueantes.append(
            "Formato sin equivalencia a galón: " + ", ".join(sin_equivalencia)
            + ". Hay que agregarlo en maestro.FORMATOS_GALONES antes de preparar."
        )
    if lec.ilegibles:
        bloqueantes.append(
            f"{_miles(lec.ilegibles)} celda(s) de colorante sin la forma CODIGO-CANTIDAD "
            f"({', '.join(lec.ejemplos_ilegibles)}). Corregirlas en el maestro."
        )
    for f, d in lec.formatos.items():
        if d["galones"] not in (None, 1):
            advertencias.append(
                f"Formato '{f}' ({_miles(d['filas'])} formulas): en Experto 2 y 3 las cantidades se dividen "
                f"por {d['galones']:g} para llevarlas a galón (Experto 1 queda tal cual)"
            )
    for p in productos.values():
        if p.repetidas:
            distintas = (
                f"; {_miles(p.distintas)} con colorantes distintos a la primera" if p.distintas
                else ", todas con los mismos colorantes"
            )
            advertencias.append(
                f"{p.producto}: {_miles(p.repetidas)} formulas duplicadas en el maestro (misma Clasificacion, "
                f"Producto, Cartilla, Color y Base{distintas}). Se deja la primera que aparece: "
                f"{_miles(p.unicas)} de {_miles(p.filas)} filas. Detalle en el resumen, hoja Duplicadas."
            )
        if p.repetidas_e3:
            distintas = (
                f"; {_miles(p.distintas_e3)} con colorantes distintos a la primera" if p.distintas_e3
                else ", todas con los mismos colorantes"
            )
            advertencias.append(
                f"{p.producto}: {_miles(p.repetidas_e3)} formulas repiten Color y Base en otra cartilla"
                f"{distintas}. Experto 3 no lleva cartilla: ahi se deja la primera que aparece "
                f"({_miles(p.unicas - p.repetidas_e3)} formulas; en Experto 1 y 2 van todas). "
                "Detalle en el resumen, hoja Duplicadas."
            )

    # --- Experto 1 del ciclo anterior ---
    glob_e1 = cfg.expertos.get(E1, EXPERTOS[E1].default_glob)
    anterior = find_latest_expert(input_dir, glob_e1)
    ant: _Lectura | None = None
    if anterior is not None and anterior.resolve() != path.resolve():
        _emit(on_progress, fase="etapa", etapa=f"Leyendo el experto anterior ({anterior.name})", pct=40)
        try:
            ant = _leer(anterior, on_progress, f"Leyendo {anterior.name}", (40, 75))
        except FormatoExpertoError as exc:
            advertencias.append(f"No se pudo leer el Experto 1 anterior ({exc}): no se compara")
    if ant is None:
        if anterior is None or anterior.resolve() == path.resolve():
            advertencias.append("No hay un Experto 1 anterior en la carpeta de entrada: no hay con qué comparar")
        anterior = None
    ant_productos = ant.productos if ant else {}

    cambios: list[Cambio] = []
    sin_cambios = agregadas_tot = quitadas_tot = 0
    for clave, p in productos.items():
        if clave not in ant_productos:
            if ant is not None:
                agregadas_tot += p.unicas  # las repetidas no se entregan
            continue
        ahora, antes = lec.contenido[clave], ant.contenido[clave]
        agregadas = sum((ahora - antes).values())
        quitadas = sum((antes - ahora).values())
        if agregadas or quitadas:
            cambios.append(Cambio(p.producto, ant_productos[clave].unicas, p.unicas, agregadas, quitadas))
            agregadas_tot += agregadas
            quitadas_tot += quitadas
        else:
            sin_cambios += 1
    quitados = [
        {"producto": p.producto, "clasificacion": p.clasificacion, "filas": p.unicas}
        for clave, p in ant_productos.items()
        if clave not in productos
    ]
    quitadas_tot += sum(q["filas"] for q in quitados)

    # --- tabla de productos ---
    tiendas, filas_tabla = leer_filas_tabla(tabla_path)
    por_e1: dict[str, FilaTabla] = {}
    for ft in filas_tabla:
        for nombre in ft.nombres.get(E1, []):
            por_e1.setdefault(normalizar(nombre), ft)
    if not por_e1:
        advertencias.append(f"{cfg.productos_name} no tiene nombres en la columna '{E1}'")
    tabla_sin_maestro = sorted(
        ft.nombres[E1][0]
        for ft in filas_tabla
        if E1 in ft.nombres and not any(normalizar(n) in productos for n in ft.nombres[E1])
    )

    # --- plantillas de Experto 2/3 y como escriben los nombres ---
    _emit(on_progress, fase="etapa", etapa="Leyendo los Experto 2 y 3 anteriores (plantillas)", pct=75)
    plantillas: dict[str, Path] = {}
    pares_e3: dict[str, tuple[str, str]] = {}
    grupos_e3: dict[str, str] = {}
    bases_e3: dict[str, str] = {}
    clasif_e2: dict[str, str] = {}
    productos_e2: dict[str, str] = {}
    for label in (E2, E3):
        plantilla = find_latest_expert(input_dir, cfg.expertos.get(label, EXPERTOS[label].default_glob))
        if plantilla is None:
            advertencias.append(
                f"No hay un {label} anterior en la carpeta de entrada que sirva de plantilla "
                f"(macros y hojas auxiliares): no se genera {label}"
            )
            continue
        try:
            enc = encabezado_libro(plantilla)
            tpl = maestro.plantilla_e3(enc) if label == E3 else maestro.plantilla_e2(enc)
            faltan = tpl.faltantes(maestro.CAMPOS_E3 if label == E3 else maestro.CAMPOS_E2)
            if faltan:
                bloqueantes.append(f"{plantilla.name}: a la plantilla le faltan las columnas {', '.join(faltan)}")
                continue
            if lec.max_colorantes > tpl.max_colorantes:
                bloqueantes.append(
                    f"El maestro tiene formulas de {lec.max_colorantes} colorantes y {plantilla.name} admite "
                    f"{tpl.max_colorantes}"
                )
            if label == E3:
                conteo = contar_valores(plantilla, ("group_code", "product_code"))
                for (g, p), _n in sorted(conteo.items(), key=lambda kv: -kv[1]):
                    if g is not None and p is not None and str(p).strip():
                        pares_e3.setdefault(normalizar(f"{g.strip()}{KEY_SEP}{p.strip()}"), (g, p))
                grupos_e3 = _canon(Counter({g: n for (g, _p), n in conteo.items() if g}))
                bases_e3 = _canon(Counter({b: n for (b,), n in contar_valores(plantilla, ("base_code",)).items()}))
            else:
                clasif_e2 = _canon(Counter({c: n for (c,), n in contar_valores(plantilla, ("clasificacion",)).items()}))
                productos_e2 = _canon(Counter({p: n for (p,), n in contar_valores(plantilla, ("producto",)).items()}))
            plantillas[label] = plantilla
        except (FormatoExpertoError, OSError, KeyError, zipfile.BadZipFile) as exc:
            bloqueantes.append(f"No se pudo leer la plantilla {plantilla.name}: {exc}")

    # --- nombres de cada producto en Experto 2 y 3 ---
    # grupo de Experto 3 de cada clasificacion del maestro (para productos sin nombre)
    grupo_votos: dict[str, Counter] = {}
    for clave, p in productos.items():
        ft = por_e1.get(clave)
        if ft is not None and E3 in ft.nombres:
            grupo = _nombre_e3(ft.nombres[E3][0], pares_e3, grupos_e3)[0]
            if grupo.strip():
                grupo_votos.setdefault(normalizar(p.clasificacion), Counter())[grupo] += 1
    grupo_de = {c: v.most_common(1)[0][0] for c, v in grupo_votos.items()}

    def grupo_por_defecto(p: ProductoMaestro) -> str:
        c = normalizar(p.clasificacion)
        return grupo_de.get(c) or grupos_e3.get(c) or p.clasificacion

    nombres: dict[str, maestro.NombresProducto] = {}
    completar: dict[int, dict[str, str]] = {}
    nuevos: list[ProductoNuevo] = []
    for clave, p in productos.items():
        ft = por_e1.get(clave)
        if ft is None:
            grupo = grupo_por_defecto(p).strip()
            nuevos.append(
                ProductoNuevo(
                    clave=clave,
                    producto=p.producto,
                    clasificacion=p.clasificacion,
                    linea=grupo,
                    filas=p.unicas,
                    repetidas=p.repetidas,
                    formatos=dict(p.formatos.most_common()),
                    cartillas=dict(p.cartillas.most_common()),
                    nombre_e2=p.producto,
                    nombre_e3=f"{grupo}{KEY_SEP}{p.producto}",
                    tiendas=list(tiendas),
                    nuevo_en_maestro=ant is not None and clave not in ant_productos,
                )
            )
            continue
        asignar: dict[str, str] = {}
        if E2 in ft.nombres:
            e2 = productos_e2.get(normalizar(ft.nombres[E2][0]), ft.nombres[E2][0])
        else:
            e2 = p.producto
            asignar[E2] = e2
        grupo, producto = _nombre_e3(ft.nombres[E3][0], pares_e3, grupos_e3) if E3 in ft.nombres else ("", "")
        if not producto:
            grupo, producto = grupo_por_defecto(p), p.producto
            asignar[E3] = f"{grupo.strip()}{KEY_SEP}{producto}"
        elif not grupo:
            grupo = grupo_por_defecto(p)
        nombres[clave] = maestro.NombresProducto(e2=e2, e3_grupo=grupo, e3_producto=producto)
        if asignar:
            completar.setdefault(ft.fila, {}).update(asignar)

    # --- que cambia en cada experto ---
    planes = [
        PlanExperto(
            label=E1,
            archivo=nombre_experto(glob_e1, fecha, path.suffix.lower()),
            plantilla=None,
            productos=len(productos),
            filas=unicas,
            entran=sorted(p.producto for c, p in productos.items() if ant is not None and c not in ant_productos),
            salen=sorted(q["producto"] for q in quitados),
            estado="ok",
        )
    ]
    por_nuevo = {n.clave: n for n in nuevos}
    for label in (E2, E3):
        plantilla = plantillas.get(label)
        glob = cfg.expertos.get(label, EXPERTOS[label].default_glob)
        if plantilla is None:
            planes.append(PlanExperto(label, None, None, 0, 0, [], [], "sin-plantilla"))
            continue
        if label == E2:
            anteriores = dict(productos_e2)
            generados = {c: normalizar(nombres[c].e2 if c in nombres else por_nuevo[c].nombre_e2) for c in productos}
        else:
            anteriores = {k: f"{g.strip()}{KEY_SEP}{p.strip()}" for k, (g, p) in pares_e3.items()}
            generados = {
                c: normalizar(
                    f"{nombres[c].e3_grupo.strip()}{KEY_SEP}{nombres[c].e3_producto}" if c in nombres
                    else por_nuevo[c].nombre_e3
                )
                for c in productos
            }
        planes.append(
            PlanExperto(
                label=label,
                archivo=nombre_experto(glob, fecha, plantilla.suffix.lower()),
                plantilla=plantilla.name,
                productos=len(productos),
                filas=unicas - repetidas_e3 if label == E3 else unicas,
                entran=sorted(productos[c].producto for c, n in generados.items() if n not in anteriores),
                salen=sorted(anteriores[k] for k in set(anteriores) - set(generados.values())),
                estado="ok",
            )
        )

    _emit(on_progress, fase="etapa", etapa="Análisis terminado", pct=100)
    return Analisis(
        maestro=path.name,
        ruta_maestro=str(path),
        fecha=fecha,
        fecha_del_nombre=del_nombre,
        filas=filas,
        filas_unicas=unicas,
        repetidas=repetidas,
        productos=len(productos),
        anterior=anterior.name if anterior else None,
        anterior_filas=sum(p.filas for p in ant_productos.values()),
        anterior_productos=len(ant_productos),
        nuevos=sorted(nuevos, key=lambda n: (n.linea, n.producto)),
        quitados=quitados,
        cambios=sorted(cambios, key=lambda c: c.producto),
        sin_cambios=sin_cambios,
        filas_agregadas=agregadas_tot,
        filas_quitadas=quitadas_tot,
        tabla_sin_maestro=tabla_sin_maestro,
        expertos=planes,
        formatos=lec.formatos,
        tiendas=tiendas,
        advertencias=advertencias,
        bloqueantes=bloqueantes,
        puede_preparar=not bloqueantes,
        repetidas_e3=repetidas_e3,
        _nombres=nombres,
        _completar=completar,
        _plantillas=plantillas,
        _pares_e3=pares_e3,
        _grupos_e3=grupos_e3,
        _bases_e3=bases_e3,
        _clasif_e2=clasif_e2,
        _productos_e2=productos_e2,
        _grupo_de=grupo_de,
        _duplicados=lec.duplicados,
    )


# --------------------------------------------------------------------------- #
# preparacion
# --------------------------------------------------------------------------- #
@dataclass
class DecisionNuevo:
    """Lo que el usuario decide para un producto nuevo."""

    tiendas: list[str]
    nombre_e2: str
    nombre_e3: str
    linea: str | None = None


@dataclass
class ArchivoGenerado:
    label: str
    archivo: str
    ruta: str
    filas: int


@dataclass
class ResultadoPreparacion:
    fecha: str
    archivos: list[ArchivoGenerado]
    tabla: str
    productos_agregados: list[dict]
    nombres_completados: int
    backup: str | None
    resumen_xlsx: str
    resumen_json: str
    advertencias: list[str]

    def to_dict(self) -> dict:
        return asdict(self)


def _filas_convertidas(
    analisis: Analisis,
    nombres: Mapping[str, maestro.NombresProducto],
    label: str,
    tpl: maestro.Plantilla,
) -> Iterator[list]:
    cols, filas = maestro.leer(Path(analisis.ruta_maestro))
    # Experto 3 no lleva cartilla: el mismo color en otra cartilla tambien se salta
    campos = maestro.CLAVE_FORMULA_E3 if label == E3 else maestro.CLAVE_FORMULA
    vistas: set[tuple] = set()
    for fila in filas:
        nombre = cols.valor(fila, "producto")
        if nombre is None or not str(nombre).strip():
            continue
        formula = cols.clave_formula(fila, campos)
        if formula in vistas:
            continue
        vistas.add(formula)
        n = nombres[normalizar(nombre)]
        colorantes = maestro.colorantes_por_galon(
            cols.celdas_colorante(fila), maestro.galones(cols.valor(fila, "formato"))
        )
        if label == E3:
            base = cols.valor(fila, "base")
            base = analisis._bases_e3.get(normalizar(base), str(base).strip()) if base else None
            yield maestro.fila_e3(cols, fila, n, base, colorantes, tpl)
        else:
            clasif = cols.valor(fila, "clasificacion")
            clasif = analisis._clasif_e2.get(normalizar(clasif), clasif) if clasif else None
            yield maestro.fila_e2(cols, fila, n, clasif, colorantes, tpl)


def _escribible(path: Path) -> bool:
    """False si el archivo esta abierto con bloqueo (Excel en Windows)."""
    try:
        with open(path, "r+b"):
            return True
    except PermissionError:
        return False
    except OSError:
        return True


def _carpeta_unica(base: Path) -> Path:
    destino, n = base, 2
    while destino.exists():
        destino = base.with_name(f"{base.name}_{n}")
        n += 1
    return destino


def preparar_expertos(
    analisis: Analisis,
    cfg: AppConfig,
    decisiones: Mapping[str, DecisionNuevo] | None = None,
    *,
    on_progress: ProgressCallback | None = None,
    ahora: datetime.datetime | None = None,
) -> ResultadoPreparacion:
    """Genera los 3 expertos del ciclo y actualiza la tabla de productos (ver
    docstring del modulo). `decisiones`: clave del producto nuevo -> lo que eligio
    el usuario; los que no esten usan lo sugerido por el analisis."""
    if not analisis.puede_preparar:
        raise PreparacionError("; ".join(analisis.bloqueantes) or "El análisis no permite preparar")
    ahora = ahora or datetime.datetime.now()
    input_dir = Path(cfg.input_dir)
    tabla = input_dir / cfg.productos_name
    staging = input_dir / STAGING_DIRNAME
    shutil.rmtree(staging, ignore_errors=True)
    staging.mkdir(parents=True)
    advertencias: list[str] = []
    nota = f"Preparación {analisis.fecha.replace('_', '/')}"

    try:
        # --- nombres de los productos nuevos (lo que decidio el usuario) ---
        nombres = dict(analisis._nombres)
        agregados: list[ProductoAgregado] = []
        for nuevo in analisis.nuevos:
            d = (decisiones or {}).get(nuevo.clave) or DecisionNuevo(nuevo.tiendas, nuevo.nombre_e2, nuevo.nombre_e3)
            e2 = (d.nombre_e2 or nuevo.producto).strip()
            grupo, producto = _nombre_e3(d.nombre_e3 or nuevo.nombre_e3, analisis._pares_e3, analisis._grupos_e3)
            if not grupo:
                grupo = (d.linea or nuevo.linea).strip()
            producto = producto or nuevo.producto
            nombres[nuevo.clave] = maestro.NombresProducto(e2=e2, e3_grupo=grupo, e3_producto=producto)
            tiendas = {t for t in d.tiendas if t in analisis.tiendas}
            agregados.append(
                ProductoAgregado(
                    linea=(d.linea or nuevo.linea or grupo).strip(),
                    producto=nuevo.producto,
                    nombres={E1: nuevo.producto, E2: e2, E3: f"{grupo.strip()}{KEY_SEP}{producto}"},
                    tiendas=tiendas,
                )
            )
            if not tiendas:
                advertencias.append(
                    f"{nuevo.producto}: se agrega a la tabla sin tiendas marcadas (no se entrega a ninguna)"
                )

        # --- Experto 1: el maestro ---
        generados: dict[str, tuple[Path, int]] = {}
        planes = {p.label: p for p in analisis.expertos}
        maestro_path = Path(analisis.ruta_maestro)
        destino_e1 = staging / planes[E1].archivo
        _emit(on_progress, fase="etapa", etapa=f"Experto 1: {destino_e1.name}", pct=2)
        if analisis.repetidas:

            def prog_e1(ev: dict) -> None:
                total = ev.get("total") or 0
                frac = min(1.0, (ev.get("leidas") or 0) / total) if total else 1.0
                _emit(on_progress, fase="progreso", etapa="Experto 1: quitando fórmulas duplicadas", pct=2 + int(frac * 13))

            _, quitadas = quitar_filas_repetidas(
                maestro_path,
                destino_e1,
                clave=[maestro.COLUMNAS[c] for c in maestro.CLAVE_FORMULA],
                on_progress=prog_e1 if on_progress else None,
            )
            advertencias.append(
                f"Experto 1: se quitaron {_miles(quitadas)} formulas duplicadas del maestro "
                "(misma Clasificacion, Producto, Cartilla, Color y Base)"
            )
        else:
            shutil.copyfile(maestro_path, destino_e1)
        generados[E1] = (destino_e1, analisis.filas_unicas)

        # --- Experto 3 y Experto 2 sobre sus plantillas ---
        tramos = {E3: (15, 55), E2: (55, 92)}
        for label in (E3, E2):
            plantilla = analisis._plantillas.get(label)
            if plantilla is None or planes[label].archivo is None:
                continue
            lo, hi = tramos[label]
            destino = staging / planes[label].archivo
            enc = encabezado_libro(plantilla)
            tpl = maestro.plantilla_e3(enc) if label == E3 else maestro.plantilla_e2(enc)
            _emit(on_progress, fase="etapa", etapa=f"{label}: {destino.name} (plantilla {plantilla.name})", pct=lo)

            def prog(ev: dict, label=label, lo=lo, hi=hi) -> None:
                if ev.get("guardando"):
                    _emit(on_progress, fase="progreso", etapa=f"{label}: guardando el libro", pct=hi - 3)
                    return
                total = ev.get("total") or 0
                frac = min(1.0, (ev.get("escritas") or 0) / total) if total else 0.0
                _emit(
                    on_progress, fase="progreso", etapa=f"{label}: {_miles(ev.get('escritas') or 0)} fórmulas",
                    pct=lo + int(frac * (hi - lo - 4)),
                )

            resultado = reescribir_hoja(
                plantilla,
                destino,
                _filas_convertidas(analisis, nombres, label, tpl),
                total=planes[label].filas,
                on_progress=prog if on_progress else None,
            )
            generados[label] = (destino, resultado.filas)
            if label == E3 and analisis.repetidas_e3:
                advertencias.append(
                    f"Experto 3: quedaron fuera {_miles(analisis.repetidas_e3)} formulas con el mismo Color y Base "
                    "en otra cartilla (Experto 3 no lleva cartilla; va la primera)"
                )

        # --- tabla de productos ---
        _emit(on_progress, fase="etapa", etapa=f"Actualizando {tabla.name}", pct=93)
        tabla_nueva = staging / tabla.name
        actualizar_tabla(tabla, tabla_nueva, agregados, analisis._completar, nota)
        completados = sum(len(v) for v in analisis._completar.values())

        # --- mover lo anterior a backups y dejar lo nuevo en la entrada ---
        _emit(on_progress, fase="etapa", etapa="Respaldando el ciclo anterior", pct=96)
        if not _escribible(tabla):
            raise PreparacionError(
                f"{tabla.name} está abierto (por ejemplo en Excel): cerrarlo y volver a preparar. No se cambió nada."
            )
        nuevos_nombres = {p.name for p, _ in generados.values()}
        for nombre in nuevos_nombres:
            existente = input_dir / nombre
            if existente.exists() and not _escribible(existente):
                raise PreparacionError(
                    f"{nombre} está abierto (por ejemplo en Excel): cerrarlo y volver a preparar. No se cambió nada."
                )
        backup = _carpeta_unica(carpeta_backups(cfg.output_dir) / BACKUP_EXPERTOS_DIRNAME / ahora.strftime(FORMATO_CARPETA))
        backup.mkdir(parents=True)
        viejos = []
        for label in EXPERTOS:
            viejos.extend(matching_files(input_dir, cfg.expertos.get(label, EXPERTOS[label].default_glob)))
        movidos: list[tuple[Path, Path]] = []
        for viejo in sorted(set(viejos)):
            try:
                os.replace(viejo, backup / viejo.name)
                movidos.append((viejo, backup / viejo.name))
            except OSError as exc:
                if viejo.name in nuevos_nombres:
                    for de, a in reversed(movidos):
                        try:
                            os.replace(a, de)
                        except OSError:
                            pass
                    raise PreparacionError(
                        f"No se pudo mover {viejo.name} al backup ({exc.strerror or exc}). "
                        "Cerrarlo y volver a preparar. No se cambió nada."
                    ) from exc
                advertencias.append(
                    f"{viejo.name} no se pudo mover al backup ({exc.strerror or exc}): queda en la carpeta de "
                    "entrada, pero el ciclo usa el experto más nuevo"
                )
        shutil.copy2(tabla, backup / tabla.name)
        os.replace(tabla_nueva, tabla)
        archivos: list[ArchivoGenerado] = []
        for label in EXPERTOS:
            if label not in generados:
                continue
            origen, filas = generados[label]
            final = input_dir / origen.name
            os.replace(origen, final)
            archivos.append(ArchivoGenerado(label, final.name, str(final), filas))
    finally:
        shutil.rmtree(staging, ignore_errors=True)

    # --- resumen ---
    _emit(on_progress, fase="etapa", etapa="Guardando el resumen", pct=98)
    carpeta = Path(cfg.output_dir) / PREPARACION_DIRNAME
    carpeta.mkdir(parents=True, exist_ok=True)
    resumen_xlsx = carpeta / f"Resumen_{analisis.fecha}.xlsx"
    resumen_json = carpeta / f"Resumen_{analisis.fecha}.json"
    resultado = ResultadoPreparacion(
        fecha=analisis.fecha,
        archivos=archivos,
        tabla=str(tabla),
        productos_agregados=[
            {"linea": a.linea, "producto": a.producto, "tiendas": sorted(a.tiendas), **a.nombres}
            for a in agregados
        ],
        nombres_completados=completados,
        backup=str(backup),
        resumen_xlsx=str(resumen_xlsx),
        resumen_json=str(resumen_json),
        advertencias=advertencias,
    )
    try:
        escribir_resumen(resumen_xlsx, analisis, resultado, ahora)
        resumen_json.write_text(
            json.dumps(
                {"preparado_en": ahora.isoformat(timespec="seconds"), "analisis": analisis.to_dict(),
                 "resultado": resultado.to_dict()},
                ensure_ascii=False,
                indent=1,
            ),
            encoding="utf-8",
        )
    except OSError as exc:
        resultado.advertencias.append(f"No se pudo guardar el resumen en {carpeta} ({exc})")
    _emit(on_progress, fase="etapa", etapa="Listo", pct=100)
    return resultado


def ultimo_resumen(cfg: AppConfig) -> dict | None:
    """El JSON de la preparacion mas reciente (por fecha de modificacion)."""
    carpeta = Path(cfg.output_dir) / PREPARACION_DIRNAME
    if not carpeta.is_dir():
        return None
    candidatos = sorted(carpeta.glob("Resumen_*.json"), key=lambda p: p.stat().st_mtime, reverse=True)
    for c in candidatos:
        try:
            return json.loads(c.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
    return None


# --------------------------------------------------------------------------- #
# resumen en Excel
# --------------------------------------------------------------------------- #
def escribir_resumen(destino: Path, analisis: Analisis, resultado: ResultadoPreparacion, ahora: datetime.datetime) -> None:
    bold = Font(bold=True)
    fondo = PatternFill("solid", fgColor="DDEBF7")
    wb = openpyxl.Workbook()

    def hoja(titulo: str, encabezado: list[str], filas: list[list], anchos: list[int]) -> None:
        ws = wb.create_sheet(titulo)
        ws.append(encabezado)
        for celda in ws[1]:
            celda.font = bold
            celda.fill = fondo
        for fila in filas:
            ws.append(fila)
        for i, ancho in enumerate(anchos):
            ws.column_dimensions[get_column_letter(i + 1)].width = ancho
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Resumen"
    datos = [
        ("Preparado el", ahora.strftime("%d/%m/%Y %H:%M")),
        ("Archivo maestro", analisis.maestro),
        ("Fecha del ciclo", analisis.fecha.replace("_", "/")),
        ("Experto 1 anterior", analisis.anterior or "-"),
        ("Fórmulas en el maestro", analisis.filas),
        ("Fórmulas duplicadas quitadas", analisis.repetidas),
        ("Fórmulas en Experto 1 y 2", analisis.filas_unicas),
        ("Mismo color en otra cartilla (no van a Experto 3)", analisis.repetidas_e3),
        ("Fórmulas en Experto 3", analisis.filas_unicas - analisis.repetidas_e3),
        ("Productos en el maestro", analisis.productos),
        ("Productos nuevos", len(analisis.nuevos)),
        ("Productos que ya no vienen", len(analisis.quitados)),
        ("Fórmulas agregadas vs anterior", analisis.filas_agregadas),
        ("Fórmulas quitadas vs anterior", analisis.filas_quitadas),
        ("Tabla de productos", resultado.tabla),
        ("Nombres de Experto 2/3 completados", resultado.nombres_completados),
        ("Backup del ciclo anterior", resultado.backup or "-"),
    ]
    datos += [(f"Generado: {a.label}", a.ruta) for a in resultado.archivos]
    for k, v in datos:
        ws.append([k, v])
    for fila in ws.iter_rows(min_col=1, max_col=1):
        fila[0].font = bold
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 90

    agregados = {normalizar(a["producto"]): a for a in resultado.productos_agregados}
    hoja(
        "Productos nuevos",
        ["Línea", "Producto", "Fórmulas", "Duplicadas quitadas", "Tiendas", "Nombre Experto 2", "Nombre Experto 3",
         "Formatos", "Cartillas", "Nuevo vs experto anterior"],
        [
            [
                n.linea, n.producto, n.filas, n.repetidas,
                ", ".join(agregados.get(n.clave, {}).get("tiendas", [])),
                agregados.get(n.clave, {}).get(E2, n.nombre_e2),
                agregados.get(n.clave, {}).get(E3, n.nombre_e3),
                ", ".join(f"{k} ({v})" for k, v in n.formatos.items()),
                ", ".join(f"{k} ({v})" for k, v in n.cartillas.items()),
                "sí" if n.nuevo_en_maestro else "no",
            ]
            for n in analisis.nuevos
        ],
        [16, 36, 10, 12, 30, 36, 40, 30, 60, 12],
    )
    hoja(
        "Cambios",
        ["Producto", "Fórmulas antes", "Fórmulas ahora", "Agregadas", "Quitadas", "Nota"],
        [[c.producto, c.antes, c.ahora, c.agregadas, c.quitadas, ""] for c in analisis.cambios]
        + [[q["producto"], q["filas"], 0, 0, q["filas"], "ya no viene en el maestro"] for q in analisis.quitados],
        [44, 14, 14, 12, 12, 30],
    )
    hoja(
        "Expertos",
        ["Experto", "Archivo", "Plantilla", "Productos", "Fórmulas", "Entran (no estaban)", "Salen"],
        [
            [p.label, p.archivo or "-", p.plantilla or "-", p.productos, p.filas, ", ".join(p.entran), ", ".join(p.salen)]
            for p in analisis.expertos
        ],
        [12, 34, 34, 10, 10, 80, 40],
    )
    hoja(
        "Duplicadas",
        ["Producto", "Clasificación", "Cartilla", "Color", "Base", "Fila que queda", "Cartilla que queda",
         "Fila quitada", "Mismos colorantes", "Se quita de"],
        [
            [d.producto, d.clasificacion, d.cartilla, d.color, d.base, d.fila_conservada, d.cartilla_conservada,
             d.fila_quitada, "sí" if d.misma_formula else "NO",
             "Experto 3 (no lleva cartilla)" if d.solo_e3 else "Expertos 1, 2 y 3"]
            for d in analisis._duplicados
        ],
        [36, 16, 24, 24, 16, 14, 24, 14, 18, 28],
    )
    hoja(
        "Advertencias",
        ["Advertencia"],
        [[a] for a in analisis.advertencias + resultado.advertencias],
        [140],
    )
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
