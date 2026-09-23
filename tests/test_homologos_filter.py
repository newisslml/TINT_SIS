import openpyxl
import pytest

from tint_sis.adapters.homologos_filter import read_expert_ids, read_homologos_ids

EXPERT_HEADER = ["ID_TINT", "group_code", "product_code", "color_key1", "R"]


@pytest.fixture
def sample_expert(tmp_path):
    path = tmp_path / "expert_idnew.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(EXPERT_HEADER)
    ws.append(["LatHab001", "Latex ", "Habitacional", "azul celeste", 10])
    ws.append(["LatHab002", "Latex ", "Habitacional", "amarillo rey", 249])
    ws.append(["EsmSCereluxe Aquatech001", "Esm. Sinteticos", "Cereluxe Aquatech", "blanco", 255])
    wb.save(path)
    return path


@pytest.fixture
def sample_homologos(tmp_path):
    # Reproduce la estructura real de las hojas de homologos ya pobladas:
    # titulo, tabla de categorias, encabezado con la columna ID_TINT en B,
    # encabezado de seccion, fila de homologo, y debajo una fila por ID_TINT.
    path = tmp_path / "homologos_idnew.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MP12"
    ws.append(["LISTADO PRODUCTOS HOMOLOGOS TWIST MP+ 12"])
    ws.append([])
    ws.append([])
    ws.append(["PRODUCTOS", "ID_TINT", "ID", "CODE", "DESCR", "PATH"])
    ws.append([None, None, 1, "Latex", "Latex", "PROD0001"])
    ws.append(["Latex -  PROD0001", None, "ID,C,5", "CODE,C,34", "DESCR,C,34", "PATH,C,12"])
    ws.append([None, None, 4, "Habitacional Ceresita", "Habitacional Ceresita", "SUBP0004"])
    ws.append([None, "LatHab001", None, "Habitacional Ceresita"])
    ws.append([None, "LatHab002", None, "Habitacional Ceresita"])
    ws.append([])
    ws.append([None, None, 11, "Cereluxe Ceresita (BS)", "Cereluxe Ceresita (BS)", "SUBP0011"])
    ws.append([None, "EsmSCereluxe Aquatech001", None, "Cereluxe Ceresita (BS)"])

    ws2 = wb.create_sheet("MP14")
    ws2.append(["PRODUCTOS", "ID_TINT", "ID", "CODE", "DESCR", "PATH"])
    ws2.append([None, None, 1, "Latex", "Latex", "PROD0001"])
    ws2.append([None, "LatHab001", None, "Habitacional Ceresita"])

    wb.save(path)
    return path


def test_read_homologos_ids_lee_la_columna_id_tint(sample_homologos):
    ids = read_homologos_ids(sample_homologos, "MP12")
    assert ids == {"LatHab001", "LatHab002", "EsmSCereluxe Aquatech001"}


def test_read_homologos_ids_ignora_filas_de_estructura(sample_homologos):
    # Numeros de categoria (col ID), encabezados de seccion ("ID,C,5"), nombres
    # de homologo y el propio encabezado "ID_TINT" no deben colarse al set.
    ids = read_homologos_ids(sample_homologos, "MP12")
    assert 1 not in ids
    assert "ID_TINT" not in ids
    assert "ID,C,5" not in ids
    assert "Habitacional Ceresita" not in ids


def test_read_homologos_ids_acepta_alias_tint_id(tmp_path):
    path = tmp_path / "homologos_alias.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MP12"
    ws.append(["PRODUCTOS", "TINT_ID", "ID", "CODE"])
    ws.append([None, "LatHab001", None, "Habitacional Ceresita"])
    ws.append([None, "LatHab002", None, "Habitacional Ceresita"])
    wb.save(path)

    assert read_homologos_ids(path, "MP12") == {"LatHab001", "LatHab002"}


def test_read_homologos_ids_sin_columna_devuelve_vacio(tmp_path):
    path = tmp_path / "homologos_sincol.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MP12"
    ws.append(["PRODUCTOS", "ID", "CODE", "DESCR"])
    ws.append([None, 4, "Habitacional Ceresita", "Habitacional Ceresita"])
    wb.save(path)

    assert read_homologos_ids(path, "MP12") == set()


def test_read_expert_ids_lee_la_columna_por_encabezado(sample_expert):
    assert read_expert_ids(sample_expert) == {"LatHab001", "LatHab002", "EsmSCereluxe Aquatech001"}
