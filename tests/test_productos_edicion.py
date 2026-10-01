import openpyxl
import pytest
from fastapi import HTTPException

from libros_prueba import HEADER_E1, HEADER_E3, experto_openpyxl
from tint_sis.adapters.productos import (
    TablaCambiadaError,
    guardar_edicion,
    leer_para_editar,
    leer_tabla,
)
from tint_sis.app import api
from tint_sis.config import AppConfig
from tint_sis.expertos import normalizar


def _tabla(path):
    """homologos_TINT.xlsx con una columna Notas y las hojas Sin asignar y Leeme."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Productos"
    ws.append(["SUBP", "Línea", "Producto", "Experto 1", "Experto 2", "Experto 3",
               "MP14", "MP12", "Tiendas 14", "Tiendas 12", "Revisar", "Notas"])
    ws.append(["SUBP0004", "Látex", "Habitacional Ceresita", "Ltx. Habitacional Ceresita",
               "Ltx.Habitacional Ceresita", "Látex / Habitacional", "x", "x", "x", "x", None, "nota 1"])
    ws.append(["SUBP0031", "Oleos", "Opaco Ceresita", "Oleo Opaco Ceresita", "Oleo Opaco Ceresita",
               "Oleos / OpacoCeresita", "x", None, None, "x", "revisar esto", None])
    ws.append([None, "Esm. al agua", "Tecnoconstrucción Mate Sipa", "E.A. Tecnoconstrucción Mate Sipa",
               "E.A. Tecnoconstrucción Mate Sipa", "Esm. al agua / Tecno Const Mate Sipa",
               None, None, None, None, None, "nota 3"])
    ws.auto_filter.ref = ws.dimensions
    wb.create_sheet("Sin asignar").append(["Experto", "Producto en el experto", "Filas"])
    wb.create_sheet("Leeme").append(["texto de ayuda"])
    wb.save(path)
    return path


def _filas(path, hoja="Productos"):
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return [list(r) for r in wb[hoja].iter_rows(values_only=True)]
    finally:
        wb.close()


def test_leer_para_editar(tmp_path):
    t = leer_para_editar(_tabla(tmp_path / "homologos_TINT.xlsx"))
    assert t["tiendas"] == ["MP14", "MP12", "Tiendas 14", "Tiendas 12"]
    assert t["version"]
    hab, opaco, mate = t["filas"]
    assert hab == {
        "id": 2, "subp": "SUBP0004", "linea": "Látex", "producto": "Habitacional Ceresita",
        "nombres": {"Experto 1": "Ltx. Habitacional Ceresita", "Experto 2": "Ltx.Habitacional Ceresita",
                    "Experto 3": "Látex / Habitacional"},
        "tiendas": ["MP14", "MP12", "Tiendas 14", "Tiendas 12"], "revisar": "",
    }
    assert opaco["tiendas"] == ["MP14", "Tiendas 12"] and opaco["revisar"] == "revisar esto"
    assert mate["tiendas"] == []


def test_guardar_edicion_marca_agrega_y_elimina(tmp_path):
    path = _tabla(tmp_path / "homologos_TINT.xlsx")
    t = leer_para_editar(path)
    hab, opaco, mate = t["filas"]
    mate["tiendas"] = ["MP14"]  # ahora va a MP14
    nuevo = {
        "id": None, "subp": "", "linea": "Texturas", "producto": "Texturex y Homologos",
        "nombres": {"Experto 1": "Texturex y Homologos", "Experto 2": "", "Experto 3": "Texturas / Texturex"},
        "tiendas": ["MP14", "Tiendas 14"],
    }
    respaldo = guardar_edicion(
        path, [hab, mate, nuevo], version=t["version"], respaldo_dir=tmp_path / "bk", nota_nuevas="Agregado en Homólogos"
    )

    filas = _filas(path)
    assert filas[0][-1] == "Notas"
    assert [f[2] for f in filas[1:]] == ["Habitacional Ceresita", "Tecnoconstrucción Mate Sipa", "Texturex y Homologos"]
    assert filas[1][11] == "nota 1"  # las notas viajan con su fila
    assert filas[2][6:10] == ["x", None, None, None] and filas[2][11] == "nota 3"
    assert filas[3] == [None, "Texturas", "Texturex y Homologos", "Texturex y Homologos", None, "Texturas / Texturex",
                        "x", None, "x", None, "Agregado en Homólogos", None]
    # Opaco se elimino; las otras hojas quedan
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Productos", "Sin asignar", "Leeme"]
    assert wb["Productos"].auto_filter.ref == "A1:L4"
    # respaldo = la tabla anterior
    assert [f[2] for f in _filas(respaldo)[1:]] == ["Habitacional Ceresita", "Opaco Ceresita", "Tecnoconstrucción Mate Sipa"]
    # el ciclo lee la tabla nueva
    tabla = leer_tabla(path)
    assert tabla.rutas["Experto 1"][normalizar("E.A. Tecnoconstrucción Mate Sipa")] == {"MP14"}
    assert normalizar("Oleo Opaco Ceresita") not in tabla.rutas["Experto 1"]


def test_guardar_edicion_no_pisa_cambios_en_disco(tmp_path):
    path = _tabla(tmp_path / "homologos_TINT.xlsx")
    t = leer_para_editar(path)
    with pytest.raises(TablaCambiadaError):
        guardar_edicion(path, t["filas"], version="123", respaldo_dir=tmp_path / "bk")
    assert not (tmp_path / "bk").exists()


@pytest.mark.parametrize(
    "cambio, error",
    [
        (lambda fs: fs[1].update(producto=" "), "sin nombre"),
        (lambda fs: fs[1]["nombres"].update({"Experto 1": "Ltx.Habitacional  Ceresita"}), "dos productos"),
        (lambda fs: fs[1].update(tiendas=["MP99"]), "MP99"),
    ],
)
def test_guardar_edicion_valida(tmp_path, cambio, error):
    path = _tabla(tmp_path / "homologos_TINT.xlsx")
    t = leer_para_editar(path)
    cambio(t["filas"])
    antes = path.read_bytes()
    with pytest.raises(ValueError, match=error):
        guardar_edicion(path, t["filas"], version=t["version"], respaldo_dir=tmp_path / "bk")
    assert path.read_bytes() == antes


def test_api_productos_y_catalogo(tmp_path, monkeypatch):
    inp = tmp_path / "data" / "input"
    inp.mkdir(parents=True)
    _tabla(inp / "homologos_TINT.xlsx")
    experto_openpyxl(
        inp / "Experto_1_28_09_2026.xlsx",
        HEADER_E1,
        [["Latex ", "Ltx. Habitacional Ceresita", "Millennium", "Galon", "amarillo", "AO-26.5"],
         ["Texturas", "Texturex y Homologos", "SW", "Tineta 4 gl", "Alamo", "OC-62.8"]],
        extra=False,
    )
    experto_openpyxl(inp / "Experto_3_28_09_2026.xlsx", HEADER_E3, [["Látex ", "Habitacional", "amarillo", "Fuerte", "AO", 26.5]])
    cfg = AppConfig(input_dir=inp, output_dir=tmp_path / "data" / "output", db_path=tmp_path / "t.db")
    monkeypatch.setattr(api, "_cfg", lambda: cfg)
    api._catalogo_cache.clear()

    t = api.get_productos()
    assert t["archivo"] == "homologos_TINT.xlsx" and len(t["filas"]) == 3

    cat = api.get_productos_catalogo()["expertos"]
    assert cat["Experto 1"]["productos"] == {"Ltx. Habitacional Ceresita": 1, "Texturex y Homologos": 1}
    assert cat["Experto 2"]["archivo"] is None
    assert cat["Experto 3"]["productos"] == {"Látex / Habitacional": 1}

    filas = t["filas"][:1]  # se eliminan Opaco y Tecnoconstruccion Mate
    r = api.put_productos({"version": t["version"], "filas": filas})
    assert r["guardado"] and len(r["filas"]) == 1
    assert "backups" in r["respaldo"] and r["respaldo"].endswith("homologos_TINT.xlsx")

    with pytest.raises(HTTPException) as exc:  # la version vieja ya no sirve
        api.put_productos({"version": t["version"], "filas": filas})
    assert exc.value.status_code == 409
