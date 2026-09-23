import hashlib
import re
import zipfile

import openpyxl
import pytest

from libros_prueba import HEADER_E3, VBA_BYTES, experto_openpyxl, experto_xlsm_a_mano
from tint_sis.adapters.sheet_filter import FormatoExpertoError, contar_claves, filtrar_libro
from tint_sis.expertos import normalizar

KEY_E3 = ("group_code", "product_code")


def _rutas(mapa: dict[str, set[str]]) -> dict[str, set[str]]:
    return {normalizar(k): v for k, v in mapa.items()}


def _filas(path, hoja="Formulas"):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return [list(r) for r in wb[hoja].iter_rows(values_only=True)]
    finally:
        wb.close()


@pytest.fixture
def e3(tmp_path):
    return experto_openpyxl(
        tmp_path / "Experto_3.xlsx",
        HEADER_E3,
        [
            ["Látex ", "Habitacional", "amarillo", "Fuerte", "AO", 26.5],
            ["Oleos", "Pajarito", "blanco", "Media", "NE", 3.5],
            ["Látex ", "Habitacional", "azul", "Media", "AZ", 9],
            ["Texturas", "Nueva", "gris", "Media", "NE", 1],
            [None, None, None, None, None, None],
            ["Látex ", "PajaritoNF", "rojo", "Media", "RJ", 1],
        ],
    )


def test_reparte_filas_por_tienda_y_las_renumera(tmp_path, e3):
    rutas = _rutas(
        {
            "Látex / Habitacional": {"T14", "T12"},
            "Oleos / Pajarito": {"T12"},
            "Látex / PajaritoNF": set(),  # conocido, no va a ninguna tienda
        }
    )
    res = filtrar_libro(e3, {"T14": tmp_path / "t14.xlsx", "T12": tmp_path / "t12.xlsx"}, rutas, KEY_E3)

    assert res.hoja == "Formulas"
    assert res.filas_por_tienda == {"T14": 2, "T12": 3}
    t14 = _filas(tmp_path / "t14.xlsx")
    assert t14[0] == HEADER_E3
    assert [f[2] for f in t14[1:]] == ["amarillo", "azul"]
    t12 = _filas(tmp_path / "t12.xlsx")
    assert [f[2] for f in t12[1:]] == ["amarillo", "blanco", "azul"]

    xml = zipfile.ZipFile(tmp_path / "t12.xlsx").read("xl/worksheets/sheet1.xml")
    assert re.findall(rb'<row r="(\d+)"', xml) == [b"1", b"2", b"3", b"4"]
    assert re.findall(rb'<c r="A(\d+)"', xml) == [b"1", b"2", b"3", b"4"]


def test_productos_fuera_de_la_tabla_se_reportan_y_no_se_entregan(tmp_path, e3):
    rutas = _rutas({"Látex / Habitacional": {"T14"}, "Látex / PajaritoNF": set()})
    res = filtrar_libro(e3, {"T14": tmp_path / "t14.xlsx"}, rutas, KEY_E3)

    # "Látex / PajaritoNF" esta en la tabla sin tiendas: no se avisa
    assert dict(res.sin_asignar) == {"Oleos / Pajarito": 1, "Texturas / Nueva": 1}
    assert res.filas_sin_clave == 1
    assert res.filas_leidas == 6


def test_la_clave_ignora_acentos_mayusculas_y_espacios(tmp_path, e3):
    rutas = _rutas({"latex/HABITACIONAL": {"T14"}})
    res = filtrar_libro(e3, {"T14": tmp_path / "t14.xlsx"}, rutas, KEY_E3)
    assert res.filas_por_tienda["T14"] == 2


def test_conserva_las_otras_hojas_y_ajusta_autofiltro(tmp_path, e3):
    rutas = _rutas({"Látex / Habitacional": {"T14"}})
    filtrar_libro(e3, {"T14": tmp_path / "t14.xlsx"}, rutas, KEY_E3)

    wb = openpyxl.load_workbook(tmp_path / "t14.xlsx")
    assert wb.sheetnames == ["Formulas", "IntegrityData"]
    assert [r for r in wb["IntegrityData"].iter_rows(values_only=True)] == [("doc_version",), (2,)]
    assert wb["Formulas"].auto_filter.ref == "A1:F3"
    xml = zipfile.ZipFile(tmp_path / "t14.xlsx").read("xl/worksheets/sheet1.xml")
    assert re.search(rb'<dimension ref="A1:F3" ?/>', xml)


def test_no_modifica_el_experto_original(tmp_path, e3):
    antes = hashlib.sha256(e3.read_bytes()).hexdigest()
    filtrar_libro(e3, {"T14": tmp_path / "t14.xlsx"}, _rutas({"Látex / Habitacional": {"T14"}}), KEY_E3)
    assert hashlib.sha256(e3.read_bytes()).hexdigest() == antes


def test_xlsm_conserva_macros_formula_del_encabezado_y_calcchain(tmp_path):
    src = experto_xlsm_a_mano(
        tmp_path / "Experto_2.xlsm",
        [
            ["Ltx. Habitacional Ceresita", "amarillo", 16.3],
            ["inline:Oleo Opaco Sipa", "blanco", 2.1],
            ["Ltx. Habitacional Ceresita", "azul", 5.5],
        ],
    )
    rutas = _rutas({"Ltx.Habitacional Ceresita": {"T14"}, "Oleo Opaco Sipa": {"T12"}})
    dst = tmp_path / "t14.xlsm"
    res = filtrar_libro(src, {"T14": dst, "T12": tmp_path / "t12.xlsm"}, rutas, ("producto",))

    assert res.filas_por_tienda == {"T14": 2, "T12": 1}
    z_src, z_dst = zipfile.ZipFile(src), zipfile.ZipFile(dst)
    assert z_dst.namelist() == z_src.namelist()
    for parte in ("xl/vbaProject.bin", "xl/worksheets/sheet2.xml", "xl/calcChain.xml", "xl/sharedStrings.xml"):
        assert z_dst.read(parte) == z_src.read(parte)
    assert z_dst.read("xl/vbaProject.bin") == VBA_BYTES

    xml = z_dst.read("xl/worksheets/sheet1.xml")
    assert b"<f>IF(1,&quot;Producto&quot;,&quot;x&quot;)</f>" in xml or b'<f>IF(1,"Producto","x")</f>' in xml
    assert b'<dimension ref="A1:C3"/>' in xml
    assert b'<autoFilter ref="A1:C3"/>' in xml
    wb_xml = z_dst.read("xl/workbook.xml")
    assert b"Formulas!$A$1:$C$3</definedName>" in wb_xml
    assert b"Colorants!$A$2:$A$65536" in wb_xml  # los demas nombres no se tocan

    filas = _filas(dst)
    assert filas[0][0] == "Producto"
    assert [f[1] for f in filas[1:]] == ["amarillo", "azul"]
    assert _filas(tmp_path / "t12.xlsm")[1] == ["Oleo Opaco Sipa", "blanco", 2.1]


def test_formulas_en_filas_de_datos_quedan_como_valor_y_se_quita_calcchain(tmp_path):
    src = experto_xlsm_a_mano(
        tmp_path / "Experto_2.xlsm",
        [
            ["Ltx. Habitacional Ceresita", ("CONCAT(\"am\",\"arillo\")", "amarillo"), 16.3],
            ["Ltx. Habitacional Ceresita", "azul", 5.5],
        ],
        calc_chain_datos=True,
    )
    dst = tmp_path / "t14.xlsm"
    filtrar_libro(src, {"T14": dst}, _rutas({"Ltx. Habitacional Ceresita": {"T14"}}), ("producto",))

    z = zipfile.ZipFile(dst)
    assert "xl/calcChain.xml" not in z.namelist()
    assert b"calcChain" not in z.read("xl/_rels/workbook.xml.rels")
    assert b"calcChain" not in z.read("[Content_Types].xml")
    xml = z.read("xl/worksheets/sheet1.xml")
    fila2 = re.search(rb'<row r="2".*?</row>', xml).group(0)
    assert b"<f>" not in fila2
    assert b't="inlineStr"' in fila2
    assert _filas(dst)[1] == ["Ltx. Habitacional Ceresita", "amarillo", 16.3]


def test_falta_columna_clave(tmp_path, e3):
    with pytest.raises(FormatoExpertoError, match="producto"):
        filtrar_libro(e3, {"T14": tmp_path / "t14.xlsx"}, {}, ("producto",))
    assert not (tmp_path / "t14.xlsx").exists()


def test_usa_la_hoja_formulas_aunque_no_sea_la_primera(tmp_path):
    path = tmp_path / "e.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "ID_new"
    wb.active.append(["otra", "cosa"])
    ws = wb.create_sheet("Formulas")
    ws.append(["Producto", "Color"])
    ws.append(["A", "rojo"])
    ws.append(["B", "verde"])
    wb.save(path)

    res = filtrar_libro(path, {"T": tmp_path / "t.xlsx"}, _rutas({"B": {"T"}}), ("producto",))
    assert res.hoja == "Formulas"
    assert _filas(tmp_path / "t.xlsx") == [["Producto", "Color"], ["B", "verde"]]


def test_emite_progreso(tmp_path, e3):
    eventos = []
    filtrar_libro(
        e3,
        {"T14": tmp_path / "t14.xlsx"},
        _rutas({"Látex / Habitacional": {"T14"}}),
        KEY_E3,
        on_progress=eventos.append,
        progress_every=2,
    )
    leidas = [e["leidas"] for e in eventos if "leidas" in e]
    assert leidas[:3] == [2, 4, 6]
    assert eventos[0]["total"] == 6  # de la dimension declarada, sin el encabezado
    assert eventos[-1] == {"guardando": True}


def test_contar_claves(e3):
    conteo = contar_claves(e3, KEY_E3)
    assert conteo == {
        "Látex / Habitacional": 2,
        "Oleos / Pajarito": 1,
        "Texturas / Nueva": 1,
        "Látex / PajaritoNF": 1,
    }
