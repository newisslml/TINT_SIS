import datetime
import json

import openpyxl
import pytest

from libros_prueba import HEADER_MAESTRO, carpeta_preparacion, experto_openpyxl
from tint_sis import preparar
from tint_sis.config import AppConfig
from tint_sis.pipeline import run_pipeline
from tint_sis.preparar import (
    DecisionNuevo,
    PreparacionError,
    analizar_maestro,
    fecha_maestro,
    nombre_experto,
    preparar_expertos,
    ultimo_resumen,
)

AHORA = datetime.datetime(2026, 9, 28, 10, 30)


@pytest.fixture
def entorno(tmp_path):
    inp, maestro = carpeta_preparacion(tmp_path / "data")
    cfg = AppConfig(input_dir=inp, output_dir=tmp_path / "data" / "output", db_path=tmp_path / "t.db")
    return cfg, maestro


def _filas(path, hoja=None):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[hoja] if hoja else wb[wb.sheetnames[0]]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


@pytest.mark.parametrize(
    "nombre, esperado",
    [
        ("Todo  MP14 28092026.xlsx", ("28_09_2026", True)),
        ("Experto_1_24_09_2026.xlsx", ("24_09_2026", True)),
        ("maestro 28-09-2026.xlsx", ("28_09_2026", True)),
        ("maestro.xlsx", ("01_10_2026", False)),
        ("Todo MP14 99992026.xlsx", ("01_10_2026", False)),
    ],
)
def test_fecha_maestro(nombre, esperado):
    assert fecha_maestro(nombre, hoy=datetime.date(2026, 10, 1)) == esperado


def test_nombre_experto():
    assert nombre_experto("Experto_1*.xlsx", "28_09_2026", ".xlsx") == "Experto_1_28_09_2026.xlsx"
    assert nombre_experto("Experto_2*.xls[xm]", "28_09_2026", ".xlsm") == "Experto_2_28_09_2026.xlsm"


def test_analisis_compara_con_el_ciclo_anterior_y_la_tabla(entorno):
    cfg, maestro = entorno
    a = analizar_maestro(maestro, cfg)

    assert (a.fecha, a.fecha_del_nombre) == ("28_09_2026", True)
    assert a.anterior == "Experto_1_24_09_2026.xlsx"
    assert (a.filas, a.filas_unicas, a.repetidas, a.productos) == (9, 7, 2, 4)
    assert a.repetidas_e3 == 1  # "amarillo" de Habitacional en otra cartilla
    assert a.puede_preparar and not a.bloqueantes

    [nuevo] = a.nuevos
    assert (nuevo.producto, nuevo.linea, nuevo.filas, nuevo.repetidas) == ("Texturex y Homologos", "Texturas", 2, 1)
    assert nuevo.nombre_e2 == "Texturex y Homologos"
    assert nuevo.nombre_e3 == "Texturas / Texturex y Homologos"
    assert nuevo.tiendas == ["MP14", "MP12", "Tiendas 14", "Tiendas 12"]
    assert nuevo.nuevo_en_maestro is True

    cambios = {c.producto: (c.antes, c.ahora, c.agregadas, c.quitadas) for c in a.cambios}
    assert cambios == {
        "Ltx. Habitacional Ceresita": (1, 2, 1, 0),
        "Ltx. CP-70 Soquina construccion": (1, 2, 1, 0),
        "Oleo Opaco Ceresita": (2, 1, 0, 1),
    }
    assert a.sin_cambios == 0
    # Habitacional +1; CP-70 +1; Texturex +2 (las duplicadas no cuentan porque no se entregan)
    assert (a.filas_agregadas, a.filas_quitadas) == (4, 1)
    assert a.quitados == []

    # duplicadas: misma Clasificacion, Producto, Cartilla, Color y Base (de los 3
    # expertos) o mismo color en otra cartilla (solo de Experto 3)
    assert [
        (d.producto, d.color, d.fila_conservada, d.fila_quitada, d.misma_formula, d.solo_e3, d.cartilla_conservada)
        for d in a._duplicados
    ] == [
        ("Texturex y Homologos", "Alamo", 6, 8, True, False, "Muestrario Texturex"),
        ("Ltx. CP-70 Soquina construccion", "Blanco", 3, 9, False, False, "Millennium"),
        ("Ltx. Habitacional Ceresita", "amarillo", 2, 10, False, True, "Millennium"),
    ]
    assert any(
        "Ltx. CP-70 Soquina construccion: 1 formulas duplicadas" in w and "1 con colorantes distintos" in w
        for w in a.advertencias
    )
    assert any(
        "Ltx. Habitacional Ceresita: 1 formulas repiten Color y Base en otra cartilla" in w for w in a.advertencias
    )

    planes = {p.label: p for p in a.expertos}
    assert planes["Experto 1"].archivo == "Experto_1_28_09_2026.xlsx"
    assert planes["Experto 1"].entran == ["Texturex y Homologos"]
    assert planes["Experto 2"].archivo == "Experto_2_28_09_2026.xlsm"
    assert planes["Experto 2"].entran == ["Ltx. CP-70 Soquina construccion", "Texturex y Homologos"]
    assert planes["Experto 3"].entran == ["Ltx. CP-70 Soquina construccion", "Texturex y Homologos"]
    assert planes["Experto 3"].salen == ["Esm. al agua / Otro"]
    # Texturex viene en Tineta 4 gl: 2 formulas (sin la repetida) pasan a galon en los 3
    assert a.filas_a_galon == 2
    assert any("'Tineta 4 gl'" in w and "divididas por 4" in w and "los 3 expertos" in w for w in a.advertencias)
    assert any("Texturex y Homologos: 1 formulas duplicadas" in w and "mismos colorantes" in w for w in a.advertencias)
    # la tabla no tiene nombre de Experto 2 para CP-70: se completa al preparar
    assert a._completar == {3: {"Experto 2": "Ltx. CP-70 Soquina construccion"}}
    json.dumps(a.to_dict())  # lo que va a la UI es serializable


def test_preparar_genera_los_tres_expertos_y_actualiza_la_tabla(entorno):
    cfg, maestro = entorno
    inp = cfg.input_dir
    a = analizar_maestro(maestro, cfg)
    r = preparar_expertos(
        a,
        cfg,
        {"texturexyhomologos": DecisionNuevo(["MP14", "Tiendas 14"], "Texturex y Homologos", "Texturas / Texturex")},
        ahora=AHORA,
    )

    assert sorted(p.name for p in inp.iterdir()) == [
        "Experto_1_28_09_2026.xlsx", "Experto_2_28_09_2026.xlsm", "Experto_3_28_09_2026.xlsx", "homologos_TINT.xlsx",
    ]
    backup = cfg.output_dir.parent / "backups" / "expertos" / "2026-09-28_10-30"
    assert r.backup == str(backup)
    assert sorted(p.name for p in backup.iterdir()) == [
        "Experto_1_24_09_2026.xlsx", "Experto_2_24_09_2026.xlsm", "Experto_3_24_09_2026.xlsx", "homologos_TINT.xlsx",
    ]
    # Experto 3 sin el "amarillo" de la otra cartilla
    assert [(x.label, x.filas) for x in r.archivos] == [("Experto 1", 7), ("Experto 2", 7), ("Experto 3", 6)]
    assert [p.filas for p in a.expertos] == [7, 7, 6]

    # Experto 1: el maestro sin las duplicadas (la segunda CP-70 "blanco" tampoco)
    e1 = _filas(inp / "Experto_1_28_09_2026.xlsx")
    assert len(e1) == 8 and e1[0][1] == "Producto "
    assert [f[6] for f in e1[1:]] == ["amarillo", "blanco", "gris", "rojo", "Alamo ", "Arcilla", "amarillo"]
    assert e1[7][2] == "Texturex"
    # ... con Texturex (Tineta 4 gl) pasado a galon; lo que ya venia en galon, igual
    assert e1[5][3] == "Galon" and e1[5][11:14] == [68.625, "OC-15.7", "VE-11.1"]
    assert e1[6][3] == "Galon" and e1[6][11:13] == [68.625, "NE-8"]
    assert e1[2][3] == "Galon (3.785 Lts.)" and e1[1][11:14] == [116, "AO-26.5", "AV-430"]
    assert any("Experto 1: 2 formulas pasadas a galón" in w for w in r.advertencias)

    # Experto 3: grafia de la plantilla, todo en galon (Texturex / 4)
    e3 = _filas(inp / "Experto_3_28_09_2026.xlsx", "Formulas")
    assert e3[1] == ["Látex ", "Habitacional", "amarillo", None, None, None, "Fuerte", "Galon", "USoz/48",
                     29.574, 48, 0, 239, 207, 30, "AO", 26.5, "AV", 430]
    assert e3[2][:9] == ["Látex ", "CP-70", "blanco", None, None, "SUGERIR USAR SOLO EN INTERIORES", "Base N",
                         "Galon", "USoz/48"]
    assert e3[2][15:17] == ["NE", 3]
    assert e3[3][6] == "Fuerte" and e3[3][15:17] == ["NE", 1.5]  # "fuerte" / "ne" normalizados
    assert e3[5][:8] == ["Texturas", "Texturex", "Alamo ", None, None, "Batir Tineta 5' antes de tintear",
                         "según producto", "Galon"]
    assert e3[5][15:19] == ["OC", 15.7, "VE", 11.1]
    assert len(e3) == 7 and [f[2] for f in e3[1:]].count("amarillo") == 1
    assert _filas(inp / "Experto_3_28_09_2026.xlsx", "IntegrityData") == [["doc_version"], [2]]

    # Experto 2: cm3, " AO", Base_Qty 3785, Note2 = Primer, clasificacion con la grafia de la plantilla
    e2 = _filas(inp / "Experto_2_28_09_2026.xlsm", "Formulas")
    assert e2[1][:10] == ["Latex ", "Ltx.Habitacional Ceresita", "Millennium", "amarillo", None, None, "Fuerte",
                          3785, " AO", pytest.approx(16.3273125)]
    assert e2[2][:2] == ["Latex ", "Ltx. CP-70 Soquina construccion"]
    assert e2[3][17] == "*USAR PRIMER GRIS"
    assert e2[5][1] == "Texturex y Homologos"
    assert e2[5][8:12] == [" OC", pytest.approx(15.7 * 29.574 / 48), " VE", pytest.approx(11.1 * 29.574 / 48)]
    assert len(e2) == 8 and e2[7][1:4] == ["Ltx.Habitacional Ceresita", "Texturex", "amarillo"]

    # tabla: producto nuevo al final y nombre de Experto 2 completado
    tabla = _filas(inp / "homologos_TINT.xlsx")
    assert tabla[2][4] == "Ltx. CP-70 Soquina construccion"
    assert tabla[2][10] == "Preparación 28/09/2026"
    assert tabla[4] == [None, "Texturas", "Texturex y Homologos", "Texturex y Homologos", "Texturex y Homologos",
                        "Texturas / Texturex", "x", None, "x", None, "Preparación 28/09/2026"]
    assert r.productos_agregados[0]["tiendas"] == ["MP14", "Tiendas 14"]

    # resumen en disco
    assert (cfg.output_dir / "Preparacion expertos" / "Resumen_28_09_2026.xlsx").exists()
    guardado = ultimo_resumen(cfg)
    assert guardado["analisis"]["nuevos"][0]["producto"] == "Texturex y Homologos"
    assert guardado["resultado"]["archivos"][0]["archivo"] == "Experto_1_28_09_2026.xlsx"
    assert not (inp / ".preparando").exists()


def test_el_ciclo_usa_lo_preparado_sin_advertencias(entorno, tmp_path):
    cfg, maestro = entorno
    preparar_expertos(analizar_maestro(maestro, cfg), cfg, ahora=AHORA)  # Texturex a todas las tiendas
    summary = run_pipeline(input_dir=cfg.input_dir, output_dir=cfg.output_dir, db_path=tmp_path / "ciclo.db")

    assert summary.ingestion_warnings == []
    filas = {(g.software, g.grupo): g.filas for g in summary.archivos}
    # Tiendas 14: Habitacional (2 cartillas en E1/E2, 1 en E3) + Texturex (2);
    # MP14: ademas CP-70 (2) y Opaco (1)
    assert filas[("Santint", "Tiendas 14")] == 3
    assert filas[("Tinwise_Lab", "Tiendas 14")] == 4
    assert filas[("Color_Pro4.8", "Tiendas 14")] == 4
    assert filas[("Corob_Tint", "MP14")] == 6
    assert filas[("Color_Pro3.1.1", "MP14")] == 7


def test_tabla_abierta_no_cambia_nada(entorno, monkeypatch):
    cfg, maestro = entorno
    antes = sorted(p.name for p in cfg.input_dir.iterdir())
    a = analizar_maestro(maestro, cfg)
    monkeypatch.setattr(preparar, "_escribible", lambda p: p.name != "homologos_TINT.xlsx")
    with pytest.raises(PreparacionError, match="abierto"):
        preparar_expertos(a, cfg, ahora=AHORA)
    assert sorted(p.name for p in cfg.input_dir.iterdir()) == antes
    assert not (cfg.output_dir.parent / "backups").exists()


def test_formato_desconocido_y_colorante_ilegible_bloquean(entorno):
    cfg, _ = entorno
    malo = experto_openpyxl(
        cfg.input_dir.parent / "maestros" / "Todo MP14 30092026.xlsx",
        HEADER_MAESTRO,
        [["Latex ", "Ltx. Habitacional Ceresita", "Millennium", "Balde", None, None, "amarillo",
          None, None, None, "Fuerte", 116, "AO26.5", None]],
        hoja="Hoja1",
        extra=False,
    )
    a = analizar_maestro(malo, cfg)
    assert not a.puede_preparar
    assert any("'Balde'" in b for b in a.bloqueantes)
    assert any("AO26.5" in b for b in a.bloqueantes)
    with pytest.raises(PreparacionError):
        preparar_expertos(a, cfg)


def test_formato_con_galones_en_el_nombre_se_convierte_y_avisa(entorno):
    cfg, _ = entorno
    balde = experto_openpyxl(
        cfg.input_dir.parent / "maestros" / "Todo MP14 30092026.xlsx",
        HEADER_MAESTRO,
        [["Latex ", "Ltx. Habitacional Ceresita", "Millennium", "Balde 5 gl", None, None, "amarillo",
          239, 207, 30, "Fuerte", 580, "AO-132.5", "AV-2150"]],
        hoja="Hoja1",
        extra=False,
    )
    a = analizar_maestro(balde, cfg)
    assert a.puede_preparar and a.filas_a_galon == 1
    assert any("'Balde 5 gl'" in w and "sale del nombre del formato" in w for w in a.advertencias)
    preparar_expertos(a, cfg, ahora=AHORA)
    e1 = _filas(cfg.input_dir / "Experto_1_30_09_2026.xlsx")
    assert e1[1][3] == "Galon" and e1[1][11:14] == [116, "AO-26.5", "AV-430"]
    e3 = _filas(cfg.input_dir / "Experto_3_30_09_2026.xlsx", "Formulas")
    assert e3[1][15:19] == ["AO", 26.5, "AV", 430]


def test_sin_plantilla_de_experto_2_no_lo_genera(entorno):
    cfg, maestro = entorno
    (cfg.input_dir / "Experto_2_24_09_2026.xlsm").unlink()
    a = analizar_maestro(maestro, cfg)
    assert a.puede_preparar
    assert any("no se genera Experto 2" in w for w in a.advertencias)
    r = preparar_expertos(a, cfg, ahora=AHORA)
    assert [x.label for x in r.archivos] == ["Experto 1", "Experto 3"]


def test_maestro_xls_no_se_acepta(entorno):
    cfg, _ = entorno
    with pytest.raises(Exception, match=".xlsx"):
        analizar_maestro(cfg.input_dir.parent / "Todo MP14 28092026.xls", cfg)
