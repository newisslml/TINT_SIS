import datetime

import openpyxl
import pytest

from tint_sis.adapters.passthrough_csv import write_passthrough_csv

HEADER = [
    "Clasificacion ", "Producto ", "Cartilla ", "Formato ", "Tolerancia luz", "Primer",
    "Color ", "R", "G", "B", "Base ", "Oz base", "Col.1", "1/48 onzas", "Col.2",
    "1/48 onzas", "Col.3", "1/48 onzas", "Col.4", "1/48 onzas", "Fecha",
]


@pytest.fixture
def sample_passthrough_excel(tmp_path):
    path = tmp_path / "expert_MP14_Corob4.1.2.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEADER)
    # fila normal: sin decimales que redondear, con blancos en RGB y Col.3/4 sin usar
    ws.append([
        "Latex ", "Ltx. Habitacional Ceresita", "Cartilla Ltx.Habitacional", "Galon",
        None, None, "azul celeste", None, None, None, "Media", 124,
        "NE", 3.5, "AZ", 9, None, None, None, None, None,
    ])
    # fila con cantidad que exige redondeo half-up (3.125 -> 3.13) y fecha 'mm-dd-yy'
    fecha_cell_row = [
        "Esmalte sintetico ", "Esm. Cereluxe", "Millennium ", "Galon", None, None,
        "CW003W", 242, 246, 237, "Blanca", 128,
        "AV", 2.5, "NE", 3.125, "OC", 1.875, None, None,
        datetime.datetime(2026, 8, 4),
    ]
    ws.append(fecha_cell_row)
    ws.cell(row=3, column=21).number_format = "mm-dd-yy"
    # fila con fecha en formato 'd/m/yyyy;@' (sin ceros de relleno)
    ws.append([
        "Latex ", "Ltx. Habitacional Ceresita", "Cartilla Ltx.Habitacional", "Galon",
        None, None, "amarillo california", 239, 207, 30, "Fuerte", 116,
        "AO", 26.5, "AV", 430, "MA", 0.5, "BC", 117,
        datetime.datetime(2026, 8, 3),
    ])
    ws.cell(row=4, column=21).number_format = "d/m/yyyy;@"
    wb.save(path)
    return path


def test_header_se_preserva_tal_cual_incluye_espacios_finales(tmp_path, sample_passthrough_excel):
    out = tmp_path / "out.csv"
    write_passthrough_csv(sample_passthrough_excel, out)
    lines = out.read_bytes().decode("iso-8859-1").split("\r\n")
    assert lines[0] == ",".join(HEADER)


def test_celdas_vacias_quedan_vacias_sin_el_relleno_de_ajuste_expert(tmp_path, sample_passthrough_excel):
    out = tmp_path / "out.csv"
    write_passthrough_csv(sample_passthrough_excel, out)
    lines = out.read_bytes().decode("iso-8859-1").split("\r\n")
    # fila "azul celeste": R,G,B vacios y Col.3/Col.4 vacios (NO "0" como en ajuste_expert)
    assert lines[1] == (
        "Latex ,Ltx. Habitacional Ceresita,Cartilla Ltx.Habitacional,Galon,,,"
        "azul celeste,,,,Media,124,NE,3.50,AZ,9.00,,,,,"
    )


def test_onzas_redondea_half_up_no_half_to_even(tmp_path, sample_passthrough_excel):
    out = tmp_path / "out.csv"
    write_passthrough_csv(sample_passthrough_excel, out)
    lines = out.read_bytes().decode("iso-8859-1").split("\r\n")
    assert "NE,3.13," in lines[2]  # 3.125 -> 3.13, no 3.12
    assert "OC,1.88," in lines[2]  # 1.875 -> 1.88


def test_fecha_respeta_el_number_format_de_la_celda(tmp_path, sample_passthrough_excel):
    out = tmp_path / "out.csv"
    write_passthrough_csv(sample_passthrough_excel, out)
    lines = out.read_bytes().decode("iso-8859-1").split("\r\n")
    assert lines[2].endswith("04/08/2026")  # 'mm-dd-yy' -> dia y mes con cero de relleno
    assert lines[3].endswith("3/8/2026")  # 'd/m/yyyy;@' -> sin ceros de relleno


def test_encoding_delimitador_y_fin_de_linea(tmp_path, sample_passthrough_excel):
    out = tmp_path / "out.csv"
    write_passthrough_csv(sample_passthrough_excel, out)
    raw = out.read_bytes()
    assert b"\r\n" in raw
    assert b'"' not in raw
    raw.decode("iso-8859-1")  # no debe lanzar


def test_devuelve_cantidad_de_filas_de_datos_escritas(tmp_path, sample_passthrough_excel):
    out = tmp_path / "out.csv"
    n = write_passthrough_csv(sample_passthrough_excel, out)
    assert n == 3
