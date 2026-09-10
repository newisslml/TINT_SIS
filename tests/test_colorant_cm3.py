import openpyxl

from tint_sis.adapters.colorant_cm3 import write_cm3_variant


def _make_src(path, *, ncols=28):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    header = [f"c{i}" for i in range(ncols)]
    header[0] = "ID_TINT"
    header[16], header[17] = "colorant_1", "qnt_ml_1"  # Q, R
    header[27] = "qnt_ml_6"  # AB -> NO se convierte
    ws.append(header)

    row = [None] * ncols
    row[0] = "LatHab001"
    row[16] = "AO"      # Q: codigo de colorante, se copia igual
    row[17] = 19.0      # R -> 19/48*29.574 = 11.706
    row[19] = 48.0      # T -> 48/48*29.574 = 29.574
    row[21] = 0.0       # V -> 0.0
    row[23] = None      # X -> queda vacio
    row[25] = 2.5       # Z -> 2.5/48*29.574 = 1.54
    row[27] = 96.0      # AB (qnt_ml_6) -> se mantiene en Oz
    ws.append(row)
    wb.save(path)


def test_convierte_solo_columnas_rtvxz(tmp_path):
    src = tmp_path / "Tiendas 14_ready.xlsx"
    dst = tmp_path / "Tiendas 14_cm3.xlsx"
    _make_src(src)

    filas = write_cm3_variant(src, dst)
    assert filas == 1

    rows = list(openpyxl.load_workbook(dst).active.iter_rows(values_only=True))
    assert rows[0][17] == "qnt_ml_1"  # encabezado intacto

    data = rows[1]
    assert data[0] == "LatHab001"     # columnas ajenas intactas
    assert data[16] == "AO"
    assert data[17] == 11.706         # 19/48*29.574
    assert data[19] == 29.574         # 48/48*29.574
    assert data[21] == 0.0            # 0 -> 0
    assert data[23] is None           # vacio queda vacio
    assert data[25] == 1.54           # 2.5/48*29.574 redondeado
    assert data[27] == 96.0           # qnt_ml_6 no se toca


def test_no_modifica_el_archivo_origen(tmp_path):
    src = tmp_path / "Tiendas 14_ready.xlsx"
    _make_src(src)
    write_cm3_variant(src, tmp_path / "out.xlsx")
    row = list(openpyxl.load_workbook(src).active.iter_rows(values_only=True))[1]
    assert row[17] == 19.0
    assert row[25] == 2.5
