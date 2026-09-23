"""Constructores de libros de prueba para los tests del filtro por productos."""
from __future__ import annotations

import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

import openpyxl

HEADER_E1 = ["Clasificacion ", "Producto ", "Cartilla ", "Formato ", "Color ", "Col.1-1/48 onzas"]
HEADER_E2 = ["Clasificacion", "Producto", "Cartilla", "Color", "Base", "Col1", "Qty1"]
HEADER_E3 = ["group_code", "product_code", "color_key1", "base_code", "colorant_1", "qnt_ml_1"]


def experto_openpyxl(path: Path, header: list, filas: list[list], *, hoja: str = "Formulas", extra: bool = True) -> Path:
    """Experto armado con openpyxl: hoja de formulas con autofiltro y, si
    `extra`, una hoja IntegrityData como la de Experto 3."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = hoja
    ws.append(header)
    for fila in filas:
        ws.append(fila)
    ws.auto_filter.ref = ws.dimensions
    if extra:
        integ = wb.create_sheet("IntegrityData")
        integ.append(["doc_version"])
        integ.append([2])
    wb.save(path)
    return path


_NS = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
_NS_R = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
VBA_BYTES = b"\xd0\xcf\x11\xe0 macro de prueba \x00\x01\x02"


def _celda(ref: str, valor, sst: list[str]) -> str:
    if valor is None:
        return ""
    if isinstance(valor, tuple):  # ("formula", "texto") -> formula con resultado de texto
        formula, resultado = valor
        return f'<c r="{ref}" t="str"><f>{escape(formula)}</f><v>{escape(resultado)}</v></c>'
    if isinstance(valor, (int, float)):
        return f'<c r="{ref}" s="1"><v>{valor}</v></c>'
    if valor.startswith("inline:"):
        return f'<c r="{ref}" t="inlineStr"><is><t>{escape(valor[7:])}</t></is></c>'
    if valor not in sst:
        sst.append(valor)
    return f'<c r="{ref}" t="s"><v>{sst.index(valor)}</v></c>'


def experto_xlsm_a_mano(path: Path, filas: list[list], *, calc_chain_datos: bool = False) -> Path:
    """Paquete .xlsm minimo escrito a mano, parecido a Experto 2: encabezado con
    una formula (A1) que el calcChain referencia, hoja auxiliar "Colorants",
    vbaProject.bin, nombres definidos y autofiltro. Columnas: A=Producto
    (formula en el encabezado), B=Color, C=Qty. Si `calc_chain_datos`, el
    calcChain ademas apunta a la fila 2 (para probar que se descarta cuando hay
    formulas en filas de datos)."""
    sst: list[str] = []
    header = [("IF(1,\"Producto\",\"x\")", "Producto"), "Color", "inline:Qty"]
    filas_xml = []
    for n, fila in enumerate([header, *filas], start=1):
        celdas = "".join(_celda(f"{col}{n}", v, sst) for col, v in zip("ABC", fila))
        filas_xml.append(f'<row r="{n}" spans="1:3">{celdas}</row>')
    ultima = len(filas) + 1
    sheet1 = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<worksheet {_NS} {_NS_R}>'
        f'<dimension ref="A1:C{ultima}"/><sheetData>{"".join(filas_xml)}</sheetData>'
        f'<autoFilter ref="A1:C{ultima}"/></worksheet>'
    )
    sheet2 = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<worksheet {_NS}>'
        f'<dimension ref="A1:A2"/><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>Code</t></is></c></row>'
        f'<row r="2"><c r="A2" t="inlineStr"><is><t>AO</t></is></c></row></sheetData></worksheet>'
    )
    sst_xml = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<sst {_NS} count="{len(sst)}" uniqueCount="{len(sst)}">'
        + "".join(f"<si><t>{escape(s)}</t></si>" for s in sst)
        + "</sst>"
    )
    calc = '<c r="A1" i="1"/>' + ('<c r="A2" i="1"/>' if calc_chain_datos else "")
    partes = {
        "[Content_Types].xml": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.ms-excel.sheet.macroEnabled.main+xml"/>'
            '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            '<Override PartName="/xl/worksheets/sheet2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
            '<Override PartName="/xl/calcChain.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/>'
            "</Types>"
        ),
        "_rels/.rels": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'<Relationship Id="rId1" Type="{_REL}/officeDocument" Target="xl/workbook.xml"/></Relationships>'
        ),
        "xl/workbook.xml": (
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<workbook {_NS} {_NS_R}><sheets>'
            '<sheet name="Formulas" sheetId="1" r:id="rId1"/><sheet name="Colorants" sheetId="2" r:id="rId2"/>'
            "</sheets><definedNames>"
            f'<definedName name="_xlnm._FilterDatabase" localSheetId="0" hidden="1">Formulas!$A$1:$C${ultima}</definedName>'
            '<definedName name="ColorantCodes">Colorants!$A$2:$A$65536</definedName>'
            "</definedNames></workbook>"
        ),
        "xl/_rels/workbook.xml.rels": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'<Relationship Id="rId1" Type="{_REL}/worksheet" Target="worksheets/sheet1.xml"/>'
            f'<Relationship Id="rId2" Type="{_REL}/worksheet" Target="worksheets/sheet2.xml"/>'
            f'<Relationship Id="rId3" Type="{_REL}/sharedStrings" Target="sharedStrings.xml"/>'
            f'<Relationship Id="rId4" Type="{_REL}/calcChain" Target="calcChain.xml"/>'
            '<Relationship Id="rId5" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" '
            'Target="vbaProject.bin"/></Relationships>'
        ),
        "xl/worksheets/sheet1.xml": sheet1,
        "xl/worksheets/sheet2.xml": sheet2,
        "xl/sharedStrings.xml": sst_xml,
        "xl/calcChain.xml": f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<calcChain {_NS}>{calc}</calcChain>',
    }
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as z:
        for nombre, contenido in partes.items():
            z.writestr(nombre, contenido)
        z.writestr("xl/vbaProject.bin", VBA_BYTES)
    return path


def tabla_productos(path: Path, filas: list[list], tiendas=("MP14", "MP12", "Tiendas 14", "Tiendas 12")) -> Path:
    """productos_TINT.xlsx: [SUBP, Linea, Producto, Experto 1, Experto 2, Experto 3, *marcas por tienda]."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Productos"
    ws.append(["SUBP", "Línea", "Producto", "Experto 1", "Experto 2", "Experto 3", *tiendas, "Revisar"])
    for fila in filas:
        ws.append(fila)
    wb.save(path)
    return path
