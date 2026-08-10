import openpyxl

from tint_sis.adapters.ajuste_expert import write_ajuste_expert_file
from tint_sis.adapters.rows import HEADER
from tint_sis.canonical.models import Colorante, FormulaCanonica


def _formula(**overrides) -> FormulaCanonica:
    data = dict(
        linea_producto="latex_experto",
        origen_archivo="x.xlsx",
        origen_fila=41,
        familia="Latex",
        clasificacion="1.-Latex",
        producto="Latex Experto",
        cartilla="C-100",
        formato="Interior",
        tolerancia_luz=None,
        primer=None,
        color="Blanco Tecnoconstrucción",
        r=255,
        g=255,
        b=255,
        base="BT",
        oz_base=100.0,
        colorantes=[Colorante(codigo="AM01", onzas_48=2.5)],
    )
    data.update(overrides)
    return FormulaCanonica(**data)


def test_writes_header_and_uses_native_excel_types(tmp_path):
    out = tmp_path / "ajuste.xlsx"
    write_ajuste_expert_file([_formula()], out)

    wb = openpyxl.load_workbook(out, data_only=True)
    ws = wb["Hoja1"]

    header_row = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header_row == HEADER

    data_row = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    assert data_row[0] == "1.-Latex"
    assert data_row[11] == 100  # Oz Base como numero, no texto
    assert data_row[12] == "AM01"
    assert data_row[13] == 2.5


def test_unused_colorante_slots_match_coroblab_quirk(tmp_path):
    out = tmp_path / "ajuste.xlsx"
    write_ajuste_expert_file([_formula(colorantes=[Colorante(codigo="AM01", onzas_48=2.5)])], out)

    wb = openpyxl.load_workbook(out, data_only=True)
    ws = wb["Hoja1"]
    data_row = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]

    # Col.2 y Col.3 sin usar muestran 0; Col.4 (el ultimo) siempre queda en blanco.
    assert data_row[14] is None  # Col. 2
    assert data_row[15] == 0  # 1/48 onzas Col. 2
    assert data_row[16] is None  # Col. 3
    assert data_row[17] == 0  # 1/48 onzas Col. 3
    assert data_row[18] is None  # Col. 4
    assert data_row[19] is None  # 1/48 onzas Col. 4
