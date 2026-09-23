import openpyxl
import pytest

from libros_prueba import HEADER_E1, HEADER_E2, HEADER_E3, experto_openpyxl, tabla_productos
from tint_sis.db import repository
from tint_sis.db.database import get_session
from tint_sis.pipeline import run_pipeline

TODAS = ["MP14", "MP12", "Tiendas 14", "Tiendas 12"]


@pytest.fixture
def entrada(tmp_path):
    """Carpeta de entrada con los 3 expertos y la tabla de productos:
      - Habitacional: todas las tiendas
      - Opaco:        solo MP14 y Tiendas 12
      - Pajarito:     en la tabla pero sin tiendas (no se entrega ni se avisa)
      - Nuevo (solo en E3): no esta en la tabla -> advertencia"""
    inp = tmp_path / "input"
    inp.mkdir()
    tabla_productos(
        inp / "productos_TINT.xlsx",
        [
            ["SUBP0004", "Látex", "Habitacional Ceresita", "Ltx. Habitacional Ceresita",
             "Ltx.Habitacional Ceresita", "Látex / Habitacional", "x", "x", "x", "x"],
            ["SUBP0031", "Oleos", "Opaco Ceresita", "Oleo Opaco Ceresita",
             "Oleo Opaco Ceresita", "Oleos / OpacoCeresita", "x", None, None, "x"],
            ["SUBP0023", "Oleos", "Pajarito Soquina", "Oleo Pajarito Soquina",
             "Oleo Pajarito Soquina", "Oleos / Pajarito", None, None, None, None],
        ],
    )
    experto_openpyxl(
        inp / "Experto_1_07_09_2026.xlsx",
        HEADER_E1,
        [
            ["Latex ", "Ltx. Habitacional Ceresita", "Millennium", "Galon", "amarillo", "AO-26.5"],
            ["Oleos", "Oleo Opaco Ceresita", "Millennium", "Galon", "blanco", "NE-3.5"],
            ["Oleos", "Oleo Pajarito Soquina", "Millennium", "Galon", "gris", "NE-1"],
            ["Latex ", "Ltx. Habitacional Ceresita", "Todocolor", "Galon", "azul", "AZ-9"],
        ],
        hoja="Todo MP14 07092026",
        extra=False,
    )
    experto_openpyxl(
        inp / "Experto_2_15_09_2026.xlsm",
        HEADER_E2,
        [
            ["Latex", "Ltx. Habitacional Ceresita", "Millennium", "amarillo", "Fuerte", " AO", 16.3273125],
            ["Oleos", "Oleo Opaco Ceresita", "Millennium", "blanco", "Media", " NE", 2.1564375],
        ],
        extra=False,
    )
    experto_openpyxl(
        inp / "Experto_3_22_09_2026.xlsx",
        HEADER_E3,
        [
            ["Látex ", "Habitacional", "amarillo", "Fuerte", "AO", 26.5],
            ["Oleos", "OpacoCeresita", "blanco", "Media", "NE", 3.5],
            ["Esm. al agua", "Tecno Const Mate Sipa", "verde", "Media", "VE", 2],
            ["Látex ", "Habitacional", "azul", "Media", "AZ", 9],
        ],
    )
    return inp


def _filas(path, hoja=None):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[hoja] if hoja else wb[wb.sheetnames[0]]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def test_entrega_a_cada_software_su_carpeta_y_formato(tmp_path, entrada):
    out = tmp_path / "output"
    summary = run_pipeline(input_dir=entrada, output_dir=out, db_path=tmp_path / "t.db")

    filtrados = out / "Archivos filtrados"
    generados = sorted(str(p.relative_to(filtrados)).replace("\\", "/") for p in filtrados.rglob("*") if p.is_file())
    assert generados == sorted(
        [
            "Santint/Tiendas 14_ready.xlsx",
            *(f"Corob_Tint/{t}_ready.xlsx" for t in TODAS),
            "Tinwise_Lab/Tiendas 14_ready.xlsm",
            *(f"Color_Pro3.1.1/{t}_ready.csv" for t in TODAS),
            "Color_Pro4.8/Tiendas 14_ready.xlsx",
            "Ibicus_Spa/Tiendas 14_ready.xlsx",
        ]
    )
    assert not (filtrados / ".staging").exists()
    assert len(summary.archivos) == 12
    assert {gf.software for gf in summary.archivos} == {
        "Santint", "Corob_Tint", "Tinwise_Lab", "Color_Pro3.1.1", "Color_Pro4.8", "Ibicus_Spa"
    }


def test_cada_tienda_recibe_solo_sus_productos(tmp_path, entrada):
    out = tmp_path / "output"
    run_pipeline(input_dir=entrada, output_dir=out, db_path=tmp_path / "t.db")
    f = out / "Archivos filtrados"

    # Experto 3 (Corob): Tiendas 12 lleva Habitacional + Opaco; Tiendas 14 solo Habitacional
    t12 = _filas(f / "Corob_Tint" / "Tiendas 12_ready.xlsx", "Formulas")
    assert t12[0] == HEADER_E3
    assert [r[2] for r in t12[1:]] == ["amarillo", "blanco", "azul"]
    assert [r[2] for r in _filas(f / "Corob_Tint" / "Tiendas 14_ready.xlsx", "Formulas")[1:]] == ["amarillo", "azul"]
    # la hoja IntegrityData de Experto 3 viaja con cada archivo
    assert openpyxl.load_workbook(f / "Santint" / "Tiendas 14_ready.xlsx").sheetnames == ["Formulas", "IntegrityData"]

    # Experto 2 (Tintwise): cantidades en cm3 tal como vienen, sin convertir
    tw = _filas(f / "Tinwise_Lab" / "Tiendas 14_ready.xlsm", "Formulas")
    assert tw[1][6] == 16.3273125 and len(tw) == 2

    # Experto 1: Color Pro 4.8 / Ibicus en Excel, Color Pro 3.1.1 en CSV (texto "AO-26.5" intacto)
    cp48 = _filas(f / "Color_Pro4.8" / "Tiendas 14_ready.xlsx")
    assert [r[4] for r in cp48[1:]] == ["amarillo", "azul"]
    csv = (f / "Color_Pro3.1.1" / "MP14_ready.csv").read_bytes().decode("iso-8859-1").split("\r\n")
    assert csv[0] == ",".join(HEADER_E1)
    assert csv[1:4] == [
        "Latex ,Ltx. Habitacional Ceresita,Millennium,Galon,amarillo,AO-26.5",
        "Oleos,Oleo Opaco Ceresita,Millennium,Galon,blanco,NE-3.5",
        "Latex ,Ltx. Habitacional Ceresita,Todocolor,Galon,azul,AZ-9",
    ]


def test_avisa_productos_que_no_estan_en_la_tabla(tmp_path, entrada):
    summary = run_pipeline(input_dir=entrada, output_dir=tmp_path / "o", db_path=tmp_path / "t.db")
    avisos = [w for w in summary.ingestion_warnings if "no estan en productos_TINT.xlsx" in w]
    assert len(avisos) == 1
    assert "Experto_3_22_09_2026.xlsx" in avisos[0] and "Tecno Const Mate Sipa (1)" in avisos[0]
    # Pajarito esta en la tabla sin tiendas: no genera aviso
    assert not any("Pajarito" in w for w in summary.ingestion_warnings)


def test_falta_un_experto_se_omiten_solo_sus_softwares(tmp_path, entrada):
    (entrada / "Experto_1_07_09_2026.xlsx").unlink()
    (entrada / "Experto_1.xls").write_bytes(b"viejo")
    out = tmp_path / "output"
    summary = run_pipeline(input_dir=entrada, output_dir=out, db_path=tmp_path / "t.db")

    softwares = {gf.software for gf in summary.archivos}
    assert softwares == {"Santint", "Corob_Tint", "Tinwise_Lab"}
    aviso = next(w for w in summary.ingestion_warnings if w.startswith("Experto 1"))
    assert ".xls" in aviso and "Color_Pro3.1.1" in aviso


def test_usa_el_experto_de_fecha_mas_nueva(tmp_path, entrada):
    experto_openpyxl(
        entrada / "Experto_3_01_08_2026.xlsx",
        HEADER_E3,
        [["Látex ", "Habitacional", "VIEJO", "Fuerte", "AO", 1]],
    )
    out = tmp_path / "output"
    run_pipeline(input_dir=entrada, output_dir=out, db_path=tmp_path / "t.db")
    colores = [r[2] for r in _filas(out / "Archivos filtrados" / "Santint" / "Tiendas 14_ready.xlsx", "Formulas")[1:]]
    assert "VIEJO" not in colores


def test_respeta_tiendas_habilitadas(tmp_path, entrada):
    out = tmp_path / "output"
    summary = run_pipeline(
        input_dir=entrada, output_dir=out, db_path=tmp_path / "t.db", enabled_grupos={"Tiendas 12"}
    )
    assert {(gf.software, gf.grupo) for gf in summary.archivos} == {
        ("Corob_Tint", "Tiendas 12"),
        ("Color_Pro3.1.1", "Tiendas 12"),
    }


def test_registra_los_archivos_en_la_base(tmp_path, entrada):
    db = tmp_path / "t.db"
    run_pipeline(input_dir=entrada, output_dir=tmp_path / "o", db_path=db)
    session = get_session(db)
    try:
        registros = repository.files_for_batch(session, repository.get_last_batch(session))
        adaptadores = {(r.linea_producto, r.adaptador) for r in registros}
    finally:
        session.close()
    assert ("Tiendas 14", "Tinwise_Lab_xlsm") in adaptadores
    assert ("MP12", "Color_Pro3.1.1_csv") in adaptadores
    assert len(registros) == 12


def test_sin_tabla_de_productos_no_genera_nada(tmp_path, entrada):
    (entrada / "productos_TINT.xlsx").unlink()
    summary = run_pipeline(input_dir=entrada, output_dir=tmp_path / "o", db_path=tmp_path / "t.db")
    assert summary.archivos == []
    assert any("tabla de productos" in w for w in summary.ingestion_warnings)


def test_eventos_de_progreso_por_experto(tmp_path, entrada):
    eventos = []
    run_pipeline(input_dir=entrada, output_dir=tmp_path / "o", db_path=tmp_path / "t.db", on_progress=eventos.append)
    fases = [e["fase"] for e in eventos]
    assert fases[0] == "inicio" and eventos[0]["total"] == 3
    assert [e["item"] for e in eventos if e["fase"] == "experto_inicio"] == ["Experto 1", "Experto 2", "Experto 3"]
    assert fases.count("experto_ok") == 3
    assert fases[-1] == "fin"


def test_expertos_desactivados_no_se_procesan_ni_avisan(tmp_path, entrada):
    eventos = []
    summary = run_pipeline(
        input_dir=entrada,
        output_dir=tmp_path / "o",
        db_path=tmp_path / "t.db",
        expertos_habilitados={"Experto 2"},
        on_progress=eventos.append,
    )
    assert {gf.software for gf in summary.archivos} == {"Tinwise_Lab"}
    assert [e["item"] for e in eventos if e["fase"] == "experto_inicio"] == ["Experto 2"]
    assert not any("Experto 1" in w or "Experto 3" in w for w in summary.ingestion_warnings)
