import openpyxl
import pytest

from libros_prueba import HEADER_E2, HEADER_E3, experto_openpyxl, tabla_productos
from tint_sis.adapters.productos import bootstrap_tabla, leer_tabla
from tint_sis.expertos import normalizar


def test_leer_tabla_rutas_por_experto_y_tienda(tmp_path):
    path = tabla_productos(
        tmp_path / "productos_TINT.xlsx",
        [
            ["SUBP0004", "Látex", "Habitacional Ceresita", "Ltx. Habitacional Ceresita",
             "Ltx. Habitacional Ceresita", "Látex / Habitacional", "x", "X", "sí", None, "revisar esto"],
            ["SUBP0043", "Látex", "Extracubriente Sipa", "Ltx. Extracubriente Sipa",
             "Ltx.Extracubriente Sipa; Ltx. Extracubriente (viejo)", "Látex / Extracubriente", "x", None, "no", "0"],
            ["SUBP0001", "Chilcomar", "Chilcomar Top 15", None, None, "Chilcomar / Top15", None, None, None, None],
        ],
    )
    tabla = leer_tabla(path)

    assert tabla.tiendas == ["MP14", "MP12", "Tiendas 14", "Tiendas 12"]
    assert tabla.productos == 3
    e2 = tabla.rutas["Experto 2"]
    assert e2[normalizar("Ltx. Habitacional Ceresita")] == {"MP14", "MP12", "Tiendas 14"}
    # varios nombres en la celda + comparacion sin espacios
    assert e2[normalizar("Ltx. Extracubriente Sipa")] == {"MP14"}
    assert e2[normalizar("Ltx. Extracubriente (viejo)")] == {"MP14"}
    # producto conocido sin tiendas: esta en la tabla con un set vacio
    assert tabla.rutas["Experto 3"][normalizar("Chilcomar / Top15")] == set()
    assert normalizar("Chilcomar / Top15") not in tabla.rutas["Experto 1"]
    assert tabla.advertencias == []


def test_leer_tabla_mismo_nombre_en_dos_filas_avisa_y_une(tmp_path):
    path = tabla_productos(
        tmp_path / "productos_TINT.xlsx",
        [
            ["S1", "Oleos", "A", None, "Oleo X", None, "x", None, None, None],
            ["S2", "Oleos", "B", None, "Oleo  X", None, None, "x", None, None],
        ],
    )
    tabla = leer_tabla(path)
    assert tabla.rutas["Experto 2"][normalizar("Oleo X")] == {"MP14", "MP12"}
    assert len(tabla.advertencias) == 1 and "Oleo" in tabla.advertencias[0]


def _homologos(path):
    """homologos_TINT.xlsx con la estructura real: tabla de lineas, secciones,
    homologos y una fila por ID_TINT (ver adapters/homologos_editor.py)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for tienda, con_oleo in (("MP14", True), ("Tiendas 12", False)):
        ws = wb.create_sheet(tienda)
        ws.append([f"LISTADO PRODUCTOS HOMOLOGOS {tienda}"])
        ws.append([])
        ws.append([])
        ws.append(["PRODUCTOS", "ID_TINT", "ID", "CODE", "DESCR", "PATH"])
        ws.append([None, None, 1, "Látex", "Látex", "PROD0001"])
        ws.append([None, None, 4, "Oleos", "Oleos", "PROD0004"])
        ws.append([])
        ws.append(["Látex -  PROD0001", None, "ID,C,5", "CODE,C,34", "DESCR,C,34", "PATH,C,12"])
        ws.append([None, None, 4, "Habitacional Ceresita", "Habitacional Ceresita", "SUBP0004"])
        ws.append([None, "LátHab001", None, "Habitacional Ceresita"])
        ws.append([None, "LátHab002", None, "Habitacional Ceresita"])
        ws.append([None, None, 9, "Nuevo Pendiente Sipa", "Nuevo Pendiente Sipa", None])
        ws.append(["Oleos -  PROD0004", None, "ID,C,5", "CODE,C,34", "DESCR,C,34", "PATH,C,12"])
        ws.append([None, None, 16, "Habitacional Ceresita", "Habitacional Ceresita", "SUBP0016"])
        if con_oleo:
            ws.append([None, "OleHab001", None, "Habitacional Ceresita"])
    wb.save(path)
    return path


def test_bootstrap_arma_la_tabla_desde_homologos_xdata_y_expertos(tmp_path):
    homologos = _homologos(tmp_path / "homologos_TINT.xlsx")
    xdata = experto_openpyxl(
        tmp_path / "xData_DATACOMPLETA_03_09_2026.xlsx",
        ["ID_TINT", "group_code", "product_code", "color_key1"],
        [
            ["LátHab001", "Látex ", "Habitacional", "amarillo"],
            ["LátHab002", "Látex ", "Habitacional", "azul"],
            ["OleHab001", "Oleos", "Habitacional", "blanco"],
        ],
    )
    e2 = experto_openpyxl(
        tmp_path / "Experto_2.xlsm",
        HEADER_E2,
        [
            ["Latex", "Ltx. Habitacional Ceresita", "Millennium", "amarillo", "Fuerte", "AO", 16.3],
            ["Oleos", "Oleo Habitacional Ceresita", "Millennium", "blanco", "Media", "NE", 2.1],
            ["Oleos", "Oleo Desconocido Sipa", "Millennium", "gris", "Media", "NE", 1.0],
        ],
        extra=False,
    )
    e3 = experto_openpyxl(
        tmp_path / "Experto_3.xlsx",
        HEADER_E3,
        [
            ["Látex ", "Habitacional", "amarillo", "Fuerte", "AO", 26.5],
            ["Oleos", "Habitacional", "blanco", "Media", "NE", 3.5],
            ["Látex ", "Nuevo Pendiente", "rojo", "Media", "RJ", 1],
        ],
    )
    destino = tmp_path / "productos_TINT.xlsx"
    resumen = bootstrap_tabla(
        homologos, xdata, {"Experto 1": None, "Experto 2": e2, "Experto 3": e3}, destino
    )

    assert resumen.productos == 3
    filas = list(openpyxl.load_workbook(destino)["Productos"].iter_rows(values_only=True))
    assert filas[0] == ("SUBP", "Línea", "Producto", "Experto 1", "Experto 2", "Experto 3",
                        "MP14", "Tiendas 12", "Revisar")
    por_nombre = {(f[1], f[2]): f for f in filas[1:]}

    lat = por_nombre[("Látex", "Habitacional Ceresita")]
    assert lat[4] == "Ltx. Habitacional Ceresita"
    assert lat[5] == "Látex / Habitacional"  # exacto via ID_TINT -> xData
    assert lat[3] == "Ltx. Habitacional Ceresita"  # Experto 1 copiado de Experto 2
    assert (lat[6], lat[7]) == ("x", "x")

    ole = por_nombre[("Oleos", "Habitacional Ceresita")]
    assert ole[4] == "Oleo Habitacional Ceresita"  # el prefijo de linea desempata
    assert ole[5] == "Oleos / Habitacional"
    assert (ole[6], ole[7]) == ("x", None)  # en Tiendas 12 el homologo no tiene IDs

    pend = por_nombre[("Látex", "Nuevo Pendiente Sipa")]
    assert (pend[6], pend[7]) == (None, None)
    assert pend[5] == "Látex / Nuevo Pendiente"  # sugerido por nombre (sin IDs)
    assert "sin formulas" in pend[8]

    sin_asignar = list(openpyxl.load_workbook(destino)["Sin asignar"].iter_rows(values_only=True))
    assert ("Experto 2", "Oleo Desconocido Sipa", 1) in sin_asignar

    # la tabla generada se lee con el mismo lector del ciclo
    tabla = leer_tabla(destino)
    assert tabla.rutas["Experto 3"][normalizar("Látex / Habitacional")] == {"MP14", "Tiendas 12"}


def test_bootstrap_no_sobreescribe(tmp_path):
    destino = tmp_path / "productos_TINT.xlsx"
    destino.write_bytes(b"revisada")
    with pytest.raises(FileExistsError):
        bootstrap_tabla(_homologos(tmp_path / "h.xlsx"), None, {}, destino)
    assert destino.read_bytes() == b"revisada"
